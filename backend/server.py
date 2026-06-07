from fastapi import FastAPI, APIRouter, HTTPException, UploadFile, File, Form, Header
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import json
import re
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Literal, Dict, Any, Union
import uuid
from datetime import datetime, timezone
from io import BytesIO

from openai import AsyncOpenAI
import auth as auth_module
from pypdf import PdfReader
from docx import Document as DocxDocument


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

print("DEBUG OPENAI KEY:", "FOUND" if os.environ.get('OPENAI_API_KEY') else "NOT FOUND")

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

auth_module.set_db(db)

async def _charge_tokens(user: dict, resp, estimated_tokens: int = 0):
    """Extract real token usage from an AI response and adjust the user's token count."""
    try:
        usage = getattr(resp, "usage", None)
        if hasattr(usage, "input_tokens") or hasattr(usage, "output_tokens"):
            actual_tokens = (usage.input_tokens or 0) + (usage.output_tokens or 0)
        else:
            actual_tokens = (
                (getattr(usage, "prompt_tokens", 0) or 0)
                + (getattr(usage, "completion_tokens", 0) or 0)
            )
        if actual_tokens > 0 and estimated_tokens > 0:
            token_diff = actual_tokens - estimated_tokens
            if token_diff != 0:
                await db.users.update_one(
                    {"id": user["id"]},
                    {"$inc": {"tokens_used_today": token_diff, "tokens_used_week": token_diff}}
                )
        elif actual_tokens > 0:
            await auth_module.check_and_charge_tokens(user, actual_tokens)
    except Exception:
        pass  # Never block the response due to tracking failure


OPENAI_API_KEY = os.environ.get('OPENAI_API_KEY', '')
FALLBACK_AI_MODEL = 'gpt-5.4-nano'
DEFAULT_AI_MODEL = os.environ.get('DEFAULT_AI_MODEL', FALLBACK_AI_MODEL)
AI_MODELS = {
    "gpt-5.4-nano": {
        "label": "GPT-5.4 nano",
        "description": "Fastest and cheapest, closest to a Haiku-style default.",
    },
    "gpt-5.4-mini": {
        "label": "GPT-5.4 mini",
        "description": "Stronger but still cost-conscious.",
    },
    "gpt-5.4": {
        "label": "GPT-5.4",
        "description": "Best quality option for harder revision tasks.",
    },
}
if DEFAULT_AI_MODEL not in AI_MODELS:
    DEFAULT_AI_MODEL = FALLBACK_AI_MODEL

IMAGE_MODELS = {
    "gpt-image-1-mini": {
        "label": "GPT Image 1 Mini",
        "description": "Fastest, lowest cost — good for quick diagrams.",
        "api_model": "gpt-image-1-mini",
        "size": "512x512",
    },
    "gpt-image-1": {
        "label": "GPT Image 1",
        "description": "Fast, low cost — good for general educational diagrams.",
        "api_model": "gpt-image-1",
        "size": "1024x1024",
    },
    "gpt-image-1.5": {
        "label": "GPT Image 1.5",
        "description": "Medium speed, better detail for complex diagrams.",
        "api_model": "gpt-image-1.5",
        "size": "1024x1024",
    },
    "gpt-image-2": {
        "label": "GPT Image 2",
        "description": "Best quality for detailed educational diagrams, higher cost.",
        "api_model": "gpt-image-2",
        "quality": "hd",
        "size": "1024x1024",
    },
}
DEFAULT_IMAGE_MODEL = "gpt-image-1"

openai_client = AsyncOpenAI(api_key=OPENAI_API_KEY) if OPENAI_API_KEY else None


def _normalise_model(model: Optional[str] = None) -> str:
    candidate = model or DEFAULT_AI_MODEL
    return candidate if candidate in AI_MODELS else DEFAULT_AI_MODEL


def _session_model(settings: Optional[dict] = None) -> str:
    return _normalise_model((settings or {}).get("model"))


def _require_ai_client():
    if not openai_client:
        raise HTTPException(status_code=500, detail="OpenAI API key not configured")


async def ai_complete(system: str, messages: List[dict], max_tokens: int, model: Optional[str] = None):
    _require_ai_client()
    return await openai_client.chat.completions.create(
        model=_normalise_model(model),
        max_completion_tokens=max_tokens,
        messages=[{"role": "system", "content": system}, *messages],
    )


def ai_text(resp) -> str:
    if not getattr(resp, "choices", None):
        return ""
    return resp.choices[0].message.content or ""


def _normalise_image_model(model: Optional[str] = None) -> str:
    candidate = model or DEFAULT_IMAGE_MODEL
    return candidate if candidate in IMAGE_MODELS else DEFAULT_IMAGE_MODEL


async def ai_image(prompt: str, model: Optional[str] = None) -> str:
    _require_ai_client()
    cfg = IMAGE_MODELS[_normalise_image_model(model)]
    kwargs: dict = {
        "model": cfg["api_model"],
        "prompt": prompt,
        "n": 1,
        "size": cfg["size"],
    }
    if cfg.get("quality"):
        kwargs["quality"] = cfg["quality"]
    resp = await openai_client.images.generate(**kwargs)
    if resp.data and resp.data[0].url:
        return resp.data[0].url
    if resp.data and resp.data[0].b64_json:
        return f"data:image/png;base64,{resp.data[0].b64_json}"
    return ""


async def ai_stream(system: str, messages: List[dict], max_tokens: int, model: Optional[str] = None):
    _require_ai_client()
    return await openai_client.chat.completions.create(
        model=_normalise_model(model),
        max_completion_tokens=max_tokens,
        messages=[{"role": "system", "content": system}, *messages],
        stream=True,
        stream_options={"include_usage": True},
    )

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


# ---------- MODELS ----------
class Subject(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    description: Optional[str] = ""
    notes: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class SubjectCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    notes: Optional[str] = ""


class SubjectUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    notes: Optional[str] = None


class ChatSessionSettings(BaseModel):
    model_config = ConfigDict(extra="ignore")
    model: str = DEFAULT_AI_MODEL
    ai_mode: Literal["normal", "quiz", "socratic", "flashcard", "exam_prep", "eli5"] = "normal"
    strictness: int = 5
    context_window: int = 0

class ChatSession(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    title: str = "New chat"
    subject_id: Optional[str] = None
    personas: List[str] = Field(default_factory=list)
    mode: Literal["solo", "group", "feynman"] = "solo"
    settings: ChatSessionSettings = Field(default_factory=ChatSessionSettings)
    system_prompt_override: Optional[str] = None
    kind: Optional[str] = None
    meta: Dict[str, Any] = Field(default_factory=dict)
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class ChatSessionCreate(BaseModel):
    title: Optional[str] = "New chat"
    subject_id: Optional[str] = None
    personas: List[str] = Field(default_factory=list)
    mode: Literal["solo", "group", "feynman"] = "solo"


class ChatSessionSettingsUpdate(BaseModel):
    model: Optional[str] = None
    ai_mode: Optional[Literal["normal", "quiz", "socratic", "flashcard", "exam_prep", "eli5"]] = None
    strictness: Optional[int] = None
    context_window: Optional[int] = None

class ChatMessage(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    session_id: str
    role: Literal["user", "assistant"]
    content: str
    persona_id: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class SendUserMessageRequest(BaseModel):
    session_id: str
    message: str


class StreamReplyRequest(BaseModel):
    session_id: str
    persona_id: Optional[str] = None


class ChatSendRequest(BaseModel):
    session_id: str
    message: str


class WorksheetRequest(BaseModel):
    subject_id: Optional[str] = None
    topic: str
    num_questions: int = 5
    difficulty: Literal["easy", "medium", "hard", "mixed"] = "medium"
    question_type: Literal["multiple_choice", "short_answer", "long_answer", "mixed"] = "mixed"
    extra_instructions: Optional[str] = ""
    model: Optional[str] = None


class WorksheetQuestion(BaseModel):
    number: Union[int, str]
    type: str
    question: str
    options: Optional[List[str]] = None
    answer: str
    explanation: Optional[str] = ""
    marks: int = 1


class MarkingFeedback(BaseModel):
    number: Union[int, str]
    awarded: float
    out_of: int
    feedback: str


class MarkingResult(BaseModel):
    total_awarded: float
    total_out_of: int
    percentage: float
    overall_feedback: str
    per_question: List[MarkingFeedback]
    marked_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class Worksheet(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    subject_id: Optional[str] = None
    subject_name: Optional[str] = ""
    topic: str
    difficulty: str
    question_type: str
    num_questions: int
    title: str
    instructions: str = ""
    total_marks: int = 0
    duration_minutes: int = 0
    questions: List[WorksheetQuestion]
    user_answers: Dict[str, str] = Field(default_factory=dict)
    marking_result: Optional[MarkingResult] = None
    confidence: Optional[Dict[str, Any]] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class MarkRequest(BaseModel):
    answers: Dict[str, str]


# ---------- AUTH / ADMIN MODELS ----------
class LoginRequest(BaseModel):
    username: str
    password: str

class RegisterRequest(BaseModel):
    username: str
    password: str

class UserLimitUpdate(BaseModel):
    token_limit_daily: Optional[int] = None
    token_limit_weekly: Optional[int] = None
    is_active: Optional[bool] = None
    is_admin: Optional[bool] = None


# ---------- HELPERS ----------
def serialize_doc(doc: dict) -> dict:
    out = dict(doc)
    for k, v in out.items():
        if isinstance(v, datetime):
            out[k] = v.isoformat()
    return out


def parse_datetime(value):
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except Exception:
            return datetime.now(timezone.utc)
    return value


async def get_subject(subject_id: str, user_id: Optional[str] = None) -> Optional[dict]:
    q = {"id": subject_id}
    if user_id:
        q["user_id"] = user_id
    return await db.subjects.find_one(q, {"_id": 0})


def build_system_message(subject: Optional[dict]) -> str:
    base = (
        "You are a patient, encouraging revision tutor for a single student. "
        "Explain concepts clearly with short paragraphs, simple examples, and use "
        "headings/bullets when helpful. Ask the student questions occasionally to "
        "check understanding. Keep answers focused and practical."
    )
    if subject:
        ctx = f"\n\nThe current revision subject is: {subject['name']}."
        if subject.get('description'):
            ctx += f"\nSubject description: {subject['description']}"
        if subject.get('notes'):
            notes = subject['notes'][:8000]
            ctx += f"\n\nReference notes provided by the student (use these as ground truth when relevant):\n---\n{notes}\n---"
        base += ctx
    return base


def extract_text_from_upload(filename: str, raw: bytes) -> str:
    name = filename.lower()
    if name.endswith('.pdf'):
        reader = PdfReader(BytesIO(raw))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(parts).strip()
    if name.endswith('.docx'):
        doc = DocxDocument(BytesIO(raw))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    try:
        return raw.decode('utf-8', errors='ignore').strip()
    except Exception:
        return ""


# ---------- BASIC ----------
@api_router.get("/")
async def root():
    return {"message": "Revisia API", "model": DEFAULT_AI_MODEL, "models": AI_MODELS}


@api_router.get("/health")
async def health():
    return {"ok": True, "has_key": bool(OPENAI_API_KEY)}


# ---------- AUTH ROUTES ----------
@api_router.post("/auth/login")
async def login(req: LoginRequest):
    return await auth_module.login_user(req.username, req.password)

@api_router.post("/auth/register")
async def register(req: RegisterRequest):
    user = await auth_module.create_user(req.username, req.password)
    return await auth_module.login_user(req.username, req.password)

@api_router.get("/auth/me")
async def me(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    return {k: v for k, v in user.items() if k != "password_hash"}


# ---------- ADMIN ROUTES ----------
@api_router.get("/admin/users")
async def admin_list_users(authorization: Optional[str] = Header(None)):
    await auth_module.require_admin(authorization)
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)
    return users

@api_router.patch("/admin/users/{user_id}")
async def admin_update_user(user_id: str, payload: UserLimitUpdate, authorization: Optional[str] = Header(None)):
    await auth_module.require_admin(authorization)
    updates = {k: v for k, v in payload.model_dump().items() if v is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="Nothing to update")
    await db.users.update_one({"id": user_id}, {"$set": updates})
    user = await db.users.find_one({"id": user_id}, {"_id": 0, "password_hash": 0})
    return user

@api_router.delete("/admin/users/{user_id}")
async def admin_delete_user(user_id: str, authorization: Optional[str] = Header(None)):
    admin = await auth_module.require_admin(authorization)
    if admin["id"] == user_id:
        raise HTTPException(status_code=400, detail="Cannot delete your own account")
    await db.users.delete_one({"id": user_id})
    return {"ok": True}

@api_router.post("/admin/users")
async def admin_create_user(req: RegisterRequest, authorization: Optional[str] = Header(None)):
    await auth_module.require_admin(authorization)
    user = await auth_module.create_user(req.username, req.password)
    return {k: v for k, v in user.items() if k != "password_hash"}

@api_router.post("/admin/users/{user_id}/reset-tokens")
async def admin_reset_tokens(user_id: str, authorization: Optional[str] = Header(None)):
    await auth_module.require_admin(authorization)
    await db.users.update_one({"id": user_id}, {"$set": {"tokens_used_today": 0, "tokens_used_week": 0}})
    return {"ok": True}


# ---------- SUBJECTS ----------
@api_router.get("/subjects", response_model=List[Subject])
async def list_subjects(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.subjects.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(500)
    for d in docs:
        d['created_at'] = parse_datetime(d.get('created_at'))
    return docs


@api_router.post("/subjects", response_model=Subject)
async def create_subject(payload: SubjectCreate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    obj = Subject(**payload.model_dump())
    doc = serialize_doc(obj.model_dump())
    doc["user_id"] = user["id"]
    await db.subjects.insert_one(doc)
    return obj


@api_router.get("/subjects/{subject_id}", response_model=Subject)
async def get_subject_endpoint(subject_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    doc = await get_subject(subject_id, user["id"])
    if not doc:
        raise HTTPException(status_code=404, detail="Subject not found")
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    return doc


@api_router.patch("/subjects/{subject_id}", response_model=Subject)
async def update_subject(subject_id: str, payload: SubjectUpdate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    update = {k: v for k, v in payload.model_dump().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="No fields to update")
    result = await db.subjects.update_one({"id": subject_id, "user_id": user["id"]}, {"$set": update})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Subject not found")
    doc = await get_subject(subject_id, user["id"])
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    return doc


@api_router.delete("/subjects/{subject_id}")
async def delete_subject(subject_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    result = await db.subjects.delete_one({"id": subject_id, "user_id": user["id"]})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Subject not found")
    await db.chat_sessions.update_many({"subject_id": subject_id, "user_id": user["id"]}, {"$set": {"subject_id": None}})
    return {"ok": True}


@api_router.post("/subjects/{subject_id}/upload")
async def upload_subject_notes(subject_id: str, file: UploadFile = File(...), append: bool = Form(True), authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    doc = await get_subject(subject_id, user["id"])
    if not doc:
        raise HTTPException(status_code=404, detail="Subject not found")
    raw = await file.read()
    text = extract_text_from_upload(file.filename or "file.txt", raw)
    if not text:
        raise HTTPException(status_code=400, detail="Could not extract text from file")
    new_notes = (doc.get('notes') or '')
    if append and new_notes:
        new_notes = new_notes + "\n\n--- " + (file.filename or 'file') + " ---\n" + text
    else:
        new_notes = text
    await db.subjects.update_one({"id": subject_id, "user_id": user["id"]}, {"$set": {"notes": new_notes}})
    return {"ok": True, "filename": file.filename, "characters": len(text)}


# ---------- CHAT ----------
@api_router.get("/chat/sessions", response_model=List[ChatSession])
async def list_sessions(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.chat_sessions.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(500)
    for d in docs:
        d['created_at'] = parse_datetime(d.get('created_at'))
    return docs


@api_router.post("/chat/sessions", response_model=ChatSession)
async def create_session(payload: ChatSessionCreate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    obj = ChatSession(**payload.model_dump())
    doc = serialize_doc(obj.model_dump())
    doc["user_id"] = user["id"]
    await db.chat_sessions.insert_one(doc)
    return obj


@api_router.delete("/chat/sessions/{session_id}")
async def delete_session(session_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    await db.chat_sessions.delete_one({"id": session_id, "user_id": user["id"]})
    await db.chat_messages.delete_many({"session_id": session_id})
    return {"ok": True}


@api_router.patch("/chat/sessions/{session_id}/settings", response_model=ChatSession)
async def update_session_settings(session_id: str, payload: ChatSessionSettingsUpdate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    session = await db.chat_sessions.find_one({"id": session_id, "user_id": user["id"]}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    current = session.get('settings') or {}
    updates = {k: v for k, v in payload.model_dump().items() if v is not None}
    if 'model' in updates:
        updates['model'] = _normalise_model(updates['model'])
    if 'strictness' in updates:
        updates['strictness'] = max(1, min(10, int(updates['strictness'])))
    if 'context_window' in updates:
        updates['context_window'] = max(0, int(updates['context_window']))
    merged = {**current, **updates}
    merged_model = ChatSessionSettings(**merged)
    await db.chat_sessions.update_one(
        {"id": session_id},
        {"$set": {"settings": merged_model.model_dump()}}
    )
    session['settings'] = merged_model.model_dump()
    session['created_at'] = parse_datetime(session.get('created_at'))
    return session


def build_mode_instructions(settings: dict) -> str:
    ai_mode = (settings or {}).get('ai_mode', 'normal')
    strictness = int((settings or {}).get('strictness', 5))
    mode_text = {
        "normal": "Be a patient, friendly tutor. Explain clearly, give examples, and check understanding occasionally.",
        "quiz": "QUIZ MODE: Take the lead. Ask the student one focused question at a time about the topic. After they answer, give a short verdict, brief explanation, and immediately ask the next question. Aim to teach them through questioning.",
        "socratic": "SOCRATIC MODE: Never give direct answers. Always respond with a guiding question that helps the student reason their way to the answer. Affirm progress, but keep the responsibility on them to think.",
        "flashcard": "FLASHCARD MODE: Each turn, give ONE flashcard in this format - 'Q: <prompt>'. Wait for the student's answer. Then reveal 'A: <answer>' and a 1-sentence note. Then give the next card. Keep cards punchy.",
        "exam_prep": "EXAM PREP MODE: Behave like an exam coach. Use exam-board language, mark schemes, command words (state, describe, explain, evaluate), and award marks out of clear totals. Always end with a quick 'next step to revise'.",
        "eli5": "ELI5 MODE: Explain like the student is 5 years old. Use everyday analogies, very short sentences, no jargon. Always include a fun, vivid example.",
    }
    strict_text = {
        1: "Be extremely lenient: accept rough answers, celebrate effort, never correct minor issues.",
        2: "Be very lenient - accept most answers, only correct major mistakes.",
        3: "Be lenient - gentle nudges on errors, focus on effort.",
        4: "Be relaxed - point out errors but don't dwell on them.",
        5: "Be balanced - fair, point out mistakes clearly but stay warm.",
        6: "Be a little firm - expect precise language.",
        7: "Be firm - demand precise language and complete reasoning.",
        8: "Be strict - call out every imprecision and missing step.",
        9: "Be very strict - exam-marker rigour, deduct for any vague phrasing.",
        10: "Be extremely strict - zero tolerance for imprecision, every word must be exam-perfect.",
    }
    return (
        f"\n\nINTERACTION MODE: {mode_text.get(ai_mode, mode_text['normal'])}"
        f"\nSTRICTNESS LEVEL ({strictness}/10): {strict_text.get(strictness, strict_text[5])}"
    )


@api_router.get("/chat/sessions/{session_id}/messages", response_model=List[ChatMessage])
async def get_messages(session_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    session = await db.chat_sessions.find_one({"id": session_id, "user_id": user["id"]}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    # Messages are scoped by session; the session was already verified to belong
    # to this user above. (Older messages were saved without a user_id field.)
    docs = await db.chat_messages.find({"session_id": session_id}, {"_id": 0}).sort("created_at", 1).to_list(2000)
    for d in docs:
        d['created_at'] = parse_datetime(d.get('created_at'))
    return docs


@api_router.post("/chat/send", response_model=ChatMessage)
async def send_message(payload: ChatSendRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()

    user = await auth_module.get_current_user(authorization)

    estimated_tokens = 2000
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    session = await db.chat_sessions.find_one({"id": payload.session_id, "user_id": user["id"]}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    model = _session_model(session.get('settings') or {})

    subject = None
    if session.get('subject_id'):
        subject = await get_subject(session['subject_id'])

    user_msg = ChatMessage(session_id=payload.session_id, role="user", content=payload.message)
    await db.chat_messages.insert_one(serialize_doc(user_msg.model_dump()))

    msg_count = await db.chat_messages.count_documents({"session_id": payload.session_id})
    if msg_count == 1 and (session.get('title') in (None, '', 'New chat')):
        title = payload.message.strip()[:60]
        await db.chat_sessions.update_one({"id": payload.session_id}, {"$set": {"title": title}})

    system_message = build_system_message(subject)
    history_docs = await db.chat_messages.find(
        {"session_id": payload.session_id}, {"_id": 0}
    ).sort("created_at", 1).to_list(2000)
    messages = [{"role": m['role'], "content": m['content']} for m in history_docs]

    try:
        resp = await ai_complete(system_message, messages, 2048, model)
        response_text = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("AI error")
        raise HTTPException(status_code=502, detail=f"AI error: {str(e)}")

    ai_msg = ChatMessage(session_id=payload.session_id, role="assistant", content=response_text)
    await db.chat_messages.insert_one(serialize_doc(ai_msg.model_dump()))
    return ai_msg


# ---------- WORKSHEETS ----------
WORKSHEET_SYSTEM = (
    "You are an expert exam-style worksheet generator. You MUST return ONLY valid JSON "
    "matching the requested schema. No prose, no markdown, no code fences."
)


def build_worksheet_prompt(req: WorksheetRequest, subject: Optional[dict]) -> str:
    type_map = {
        "multiple_choice": "Every question must be multiple choice with exactly 4 options labelled A-D.",
        "short_answer": "Every question must be a short-answer question (1-3 sentences).",
        "long_answer": "Every question must be a long-answer / extended response question.",
        "mixed": "Mix of multiple choice (with 4 options), short answer, and one or two long answer questions. Roughly 50% MCQ, 30% short, 20% long.",
    }
    parts = [
        f"Generate a revision worksheet on the topic: \"{req.topic}\".",
        f"Number of questions: {req.num_questions}.",
        f"Difficulty: {req.difficulty}.",
        f"Question style: {type_map[req.question_type]}",
    ]
    if subject:
        parts.append(f"Subject: {subject['name']}.")
        if subject.get('description'):
            parts.append(f"Subject context: {subject['description']}")
        if subject.get('notes'):
            parts.append(f"Use these student notes as ground truth where relevant:\n---\n{subject['notes'][:6000]}\n---")
    if req.extra_instructions:
        parts.append(f"Additional instructions: {req.extra_instructions}")
    parts.append(
        "Return ONLY a JSON object with this exact shape (proper exam paper structure):\n"
        "{\n"
        '  "title": "<exam-style title>",\n'
        '  "instructions": "<2-4 short bullet-style sentences>",\n'
        '  "duration_minutes": <integer>,\n'
        '  "total_marks": <integer>,\n'
        '  "questions": [\n'
        '    {\n'
        '      "number": 1,\n'
        '      "type": "multiple_choice" | "short_answer" | "long_answer",\n'
        '      "question": "<the question>",\n'
        '      "options": ["A) ...", "B) ...", "C) ...", "D) ..."],\n'
        '      "marks": <integer>,\n'
        '      "answer": "<model answer>",\n'
        '      "explanation": "<1-3 sentence markscheme notes>"\n'
        '    }\n'
        '  ]\n'
        '}\n'
        "Do not wrap in code fences. Do not include any text outside the JSON."
    )
    return "\n\n".join(parts)


def parse_worksheet_json(text: str) -> dict:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    match = re.search(r"\{[\s\S]*\}", cleaned)
    if not match:
        raise ValueError("No JSON object found in model output")
    blob = match.group(0)
    try:
        return json.loads(blob)
    except json.JSONDecodeError:
        for end in range(len(blob), 0, -1):
            candidate = blob[:end]
            try:
                return json.loads(candidate)
            except json.JSONDecodeError:
                pass
            if end < len(blob) - 800:
                break
        opens_obj = blob.count('{') - blob.count('}')
        opens_arr = blob.count('[') - blob.count(']')
        repair = blob.rstrip().rstrip(',')
        quotes = len(re.findall(r'(?<!\\)"', repair))
        if quotes % 2 == 1:
            repair += '"'
        repair += (']' * max(0, opens_arr)) + ('}' * max(0, opens_obj))
        return json.loads(repair)


PAST_PAPER_SYSTEM = (
    "You are an expert exam paper analyst. You MUST return ONLY valid JSON "
    "matching the requested schema. No prose, no markdown, no code fences."
)


def build_past_paper_prompt(text: str, subject: Optional[dict], difficulty: str, num_questions: Optional[int]) -> str:
    parts = [
        "Below is the text of a real past exam paper. Your job is to extract the questions "
        "as faithfully as possible and reformat them into the standard worksheet JSON schema.",
        "",
        "RULES:",
        "- Preserve each question's wording exactly as written. Do not rephrase or simplify.",
        "- If the paper has sections (Section A, Section B, etc.), prefix question numbers like A1, A2, B1.",
        "- Keep the original mark allocation for each question.",
        "- If a question has multiple parts (1a, 1b, etc.), flatten them into separate numbered questions.",
        "- Assign a question type ('multiple_choice', 'short_answer', or 'long_answer') based on the answer space.",
        "- For multiple-choice questions, include the 4 options as an array.",
        "- Include a model answer based on what the mark scheme would expect.",
        "- Include a brief explanation of what gains marks.",
        f"- Difficulty: {difficulty}.",
    ]
    if num_questions:
        parts.append(f"- Extract at most {num_questions} questions (prefer the first {num_questions}).")
    if subject:
        parts.append(f"- Subject: {subject['name']}.")
    parts.append(
        "",
        "Return ONLY a JSON object with this exact shape:\n"
        "{\n"
        '  "title": "<original paper title or best guess>",\n'
        '  "instructions": "<2-4 bullet-style exam instructions, or '
        'copy the rubric from the paper>",\n'
        '  "duration_minutes": <integer, 0 if unknown>,\n'
        '  "total_marks": <integer, sum of all extracted question marks>,\n'
        '  "questions": [\n'
        '    {\n'
        '      "number": "1",\n'
        '      "type": "multiple_choice" | "short_answer" | "long_answer",\n'
        '      "question": "<the question text exactly as written>",\n'
        '      "options": ["A) ...", "B) ...", "C) ...", "D) ..."],\n'
        '      "marks": <integer>,\n'
        '      "answer": "<model answer>",\n'
        '      "explanation": "<markscheme notes>"\n'
        '    }\n'
        '  ]\n'
        '}\n'
        "Do not wrap in code fences. Do not include any text outside the JSON.",
        "",
        "--- PAST PAPER TEXT ---",
        text[:15000],
    )
    return "\n".join(parts)


@api_router.post("/worksheets/from-past-paper", response_model=Worksheet)
async def worksheet_from_past_paper(
    file: UploadFile = File(...),
    subject_id: Optional[str] = Form(None),
    difficulty: str = Form("hard"),
    num_questions: Optional[int] = Form(None),
    authorization: Optional[str] = Header(None),
):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)

    estimated_tokens = 8000
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    raw = await file.read()
    text = extract_text_from_upload(file.filename or "paper.pdf", raw)
    if not text or len(text.strip()) < 50:
        raise HTTPException(
            status_code=400,
            detail="Could not extract enough text from the file. Scanned/image-only PDFs are not supported."
        )

    subject = None
    if subject_id:
        subject = await get_subject(subject_id, user["id"])

    prompt = build_past_paper_prompt(text, subject, difficulty, num_questions)
    try:
        response = await ai_complete(PAST_PAPER_SYSTEM, [{"role": "user", "content": prompt}], 4096)
        raw_json = ai_text(response)
        await _charge_tokens(user, response, estimated_tokens)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("AI error")
        raise HTTPException(status_code=502, detail=f"AI error: {str(e)}")

    try:
        data = parse_worksheet_json(raw_json)
    except Exception as e:
        logger.error(f"Past-paper parse failed. Raw output: {raw_json[:1000]}")
        raise HTTPException(status_code=502, detail=f"Could not parse worksheet JSON: {e}")

    questions = []
    for i, q in enumerate(data.get('questions', []), start=1):
        questions.append(WorksheetQuestion(
            number=str(q.get('number') or i),
            type=q.get('type', 'short_answer'),
            question=q.get('question', ''),
            options=q.get('options') if q.get('options') else None,
            answer=q.get('answer', ''),
            explanation=q.get('explanation', '') or '',
            marks=int(q.get('marks') or 1),
        ))

    total_marks = data.get('total_marks') or sum(q.marks for q in questions)
    duration = data.get('duration_minutes') or 0
    instructions = data.get('instructions') or (
        "Answer ALL questions in the spaces provided. Read each question carefully."
    )

    ws = Worksheet(
        subject_id=subject_id,
        subject_name=subject['name'] if subject else "",
        topic="Past paper",
        difficulty=difficulty,
        question_type="mixed",
        num_questions=len(questions),
        title=data.get('title') or (file.filename or "Past paper").rsplit('.', 1)[0],
        instructions=instructions,
        total_marks=int(total_marks),
        duration_minutes=int(duration),
        questions=questions,
    )
    ws_doc = serialize_doc(ws.model_dump())
    ws_doc["user_id"] = user["id"]
    await db.worksheets.insert_one(ws_doc)
    return ws


@api_router.post("/worksheets/generate", response_model=Worksheet)
async def generate_worksheet(req: WorksheetRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)

    estimated_tokens = 3500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    subject = None
    if req.subject_id:
        subject = await get_subject(req.subject_id, user["id"])

    prompt = build_worksheet_prompt(req, subject)
    try:
        response = await ai_complete(WORKSHEET_SYSTEM, [{"role": "user", "content": prompt}], 4096, model=req.model)
        raw = ai_text(response)
        await _charge_tokens(user, response, estimated_tokens)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("AI error")
        raise HTTPException(status_code=502, detail=f"AI error: {str(e)}")

    try:
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.error(f"Worksheet parse failed. Raw output: {raw[:1000]}")
        raise HTTPException(status_code=502, detail=f"Could not parse worksheet JSON: {e}")

    questions = []
    for i, q in enumerate(data.get('questions', []), start=1):
        questions.append(WorksheetQuestion(
            number=q.get('number') or i,
            type=q.get('type', 'short_answer'),
            question=q.get('question', ''),
            options=q.get('options') if q.get('options') else None,
            answer=q.get('answer', ''),
            explanation=q.get('explanation', '') or '',
            marks=int(q.get('marks') or (1 if q.get('type') == 'multiple_choice' else 3)),
        ))

    total_marks = data.get('total_marks') or sum(q.marks for q in questions)
    duration = data.get('duration_minutes') or max(10, total_marks * 1)
    instructions = data.get('instructions') or (
        "Answer ALL questions in the spaces provided. "
        "Read each question carefully. "
        "Show your working where appropriate."
    )

    ws = Worksheet(
        subject_id=req.subject_id,
        subject_name=subject['name'] if subject else "",
        topic=req.topic,
        difficulty=req.difficulty,
        question_type=req.question_type,
        num_questions=req.num_questions,
        title=data.get('title') or f"Worksheet: {req.topic}",
        instructions=instructions,
        total_marks=int(total_marks),
        duration_minutes=int(duration),
        questions=questions,
    )
    ws_doc = serialize_doc(ws.model_dump())
    ws_doc["user_id"] = user["id"]
    await db.worksheets.insert_one(ws_doc)
    return ws


@api_router.get("/worksheets", response_model=List[Worksheet])
async def list_worksheets(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.worksheets.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)
    for d in docs:
        d['created_at'] = parse_datetime(d.get('created_at'))
        if d.get('marking_result') and isinstance(d['marking_result'].get('marked_at'), str):
            d['marking_result']['marked_at'] = parse_datetime(d['marking_result']['marked_at'])
    return docs


@api_router.get("/worksheets/{worksheet_id}", response_model=Worksheet)
async def get_worksheet(worksheet_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    doc = await db.worksheets.find_one({"id": worksheet_id, "user_id": user["id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Worksheet not found")
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    if doc.get('marking_result') and isinstance(doc['marking_result'].get('marked_at'), str):
        doc['marking_result']['marked_at'] = parse_datetime(doc['marking_result']['marked_at'])
    return doc


MARKER_SYSTEM = (
    "You are a fair, encouraging exam marker. You MUST return ONLY valid JSON, no prose, no code fences."
)


@api_router.post("/worksheets/{worksheet_id}/mark", response_model=Worksheet)
async def mark_worksheet(worksheet_id: str, payload: MarkRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)

    estimated_tokens = 2500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    doc = await db.worksheets.find_one({"id": worksheet_id, "user_id": user["id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Worksheet not found")

    lines = [
        f"Mark this {doc.get('subject_name') or 'revision'} worksheet titled \"{doc['title']}\".",
        "For each question, compare the student's answer to the model answer and award marks fairly.",
        "Give partial credit. Be concise and encouraging in feedback (1-2 sentences each).",
        "",
        "Questions and student answers:",
    ]
    for q in doc['questions']:
        num = q['number']
        student = payload.answers.get(str(num), "").strip() or "[no answer]"
        lines.append(f"\nQ{num} [{q.get('marks', 1)} marks] ({q['type']}): {q['question']}")
        if q.get('options'):
            lines.append("Options: " + " | ".join(q['options']))
        lines.append(f"Model answer: {q['answer']}")
        if q.get('explanation'):
            lines.append(f"Markscheme notes: {q['explanation']}")
        lines.append(f"Student answer: {student}")

    lines.append(
        "\nReturn ONLY this JSON:\n"
        "{\n"
        '  "per_question": [\n'
        '    {"number": <int>, "awarded": <number>, "out_of": <int>, "feedback": "<1-2 sentences>"}\n'
        '  ],\n'
        '  "overall_feedback": "<2-3 sentences>"\n'
        "}"
    )
    prompt = "\n".join(lines)

    try:
        response = await ai_complete(MARKER_SYSTEM, [{"role": "user", "content": prompt}], 2048)
        raw = ai_text(response)
        await _charge_tokens(user, response, estimated_tokens)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Marker error")
        raise HTTPException(status_code=502, detail=f"AI error: {str(e)}")

    try:
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.error(f"Marking parse failed: {raw[:500]}")
        raise HTTPException(status_code=502, detail=f"Could not parse marking result: {e}")

    per_q = []
    total_awarded = 0.0
    total_out_of = 0
    by_num = {q['number']: q for q in doc['questions']}
    for item in data.get('per_question', []):
        num = int(item.get('number', 0))
        q = by_num.get(num)
        out_of = int(item.get('out_of') or (q.get('marks', 1) if q else 1))
        awarded = float(item.get('awarded') or 0)
        awarded = max(0.0, min(awarded, out_of))
        per_q.append(MarkingFeedback(
            number=num, awarded=awarded, out_of=out_of,
            feedback=item.get('feedback', '') or ''
        ))
        total_awarded += awarded
        total_out_of += out_of

    if total_out_of == 0:
        total_out_of = sum(q.get('marks', 1) for q in doc['questions'])

    result = MarkingResult(
        total_awarded=round(total_awarded, 1),
        total_out_of=total_out_of,
        percentage=round((total_awarded / total_out_of) * 100, 1) if total_out_of else 0.0,
        overall_feedback=data.get('overall_feedback', '') or '',
        per_question=per_q,
    )

    result_doc = serialize_doc(result.model_dump())
    await db.worksheets.update_one(
        {"id": worksheet_id},
        {"$set": {"user_answers": payload.answers, "marking_result": result_doc}}
    )

    doc['user_answers'] = payload.answers
    doc['marking_result'] = result_doc
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    if isinstance(doc['marking_result'].get('marked_at'), str):
        doc['marking_result']['marked_at'] = parse_datetime(doc['marking_result']['marked_at'])
    return doc


@api_router.delete("/worksheets/{worksheet_id}")
async def delete_worksheet(worksheet_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    res = await db.worksheets.delete_one({"id": worksheet_id, "user_id": user["id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Worksheet not found")
    return {"ok": True}


# ---------- PERSONAS ----------
PERSONAS = {
    "einstein": {
        "id": "einstein", "name": "Albert Einstein", "title": "Theoretical physicist",
        "era": "1879-1955", "tags": ["physics", "relativity", "mathematics"],
        "system_prompt": (
            "You are Albert Einstein. Speak in first person as Einstein would: thoughtful, playful, "
            "fond of vivid thought experiments and gentle humour with a faint German cadence. "
            "Reference relativity, quantum debates, your time at the patent office, Princeton, and your "
            "love of music when relevant. Stay in character even if asked modern questions. "
            "Keep answers warm and accessible."
        ),
    },
    "newton": {
        "id": "newton", "name": "Isaac Newton", "title": "Mathematician & physicist",
        "era": "1643-1727", "tags": ["physics", "mathematics", "calculus", "gravity"],
        "system_prompt": (
            "You are Sir Isaac Newton. Be formal, meticulous, occasionally aloof. Reference your laws of "
            "motion, gravitation, the Principia, your time at Cambridge, your work on optics and calculus, "
            "and your alchemical curiosities. Use a slightly archaic, precise English."
        ),
    },
    "curie": {
        "id": "curie", "name": "Marie Curie", "title": "Physicist & chemist",
        "era": "1867-1934", "tags": ["chemistry", "physics", "radioactivity"],
        "system_prompt": (
            "You are Marie Curie. Quiet, methodical, fiercely determined. Refer to your work isolating "
            "polonium and radium, your Nobel prizes, the radium institute, the X-ray ambulances in WWI, "
            "and the obstacles you faced as a woman in science. Be encouraging to learners."
        ),
    },
    "darwin": {
        "id": "darwin", "name": "Charles Darwin", "title": "Naturalist",
        "era": "1809-1882", "tags": ["biology", "evolution", "ecology"],
        "system_prompt": (
            "You are Charles Darwin. Patient, observant, slightly hesitant scholar. Reference the Beagle "
            "voyage, the Galapagos finches, natural selection, your decades of caution before publishing "
            "On the Origin of Species. Use careful, deliberate Victorian prose."
        ),
    },
    "davinci": {
        "id": "davinci", "name": "Leonardo da Vinci", "title": "Polymath",
        "era": "1452-1519", "tags": ["art", "anatomy", "engineering", "design"],
        "system_prompt": (
            "You are Leonardo da Vinci. Endlessly curious, sketch metaphors into every explanation, "
            "blend art and engineering. Refer to your notebooks, anatomy studies, flying machines, "
            "and Florentine workshops. Speak with wonder."
        ),
    },
    "shakespeare": {
        "id": "shakespeare", "name": "William Shakespeare", "title": "Playwright",
        "era": "1564-1616", "tags": ["literature", "drama", "poetry", "english"],
        "system_prompt": (
            "You are William Shakespeare. Speak with theatrical flair, slip into iambic pentameter "
            "when it serves, quote your own works freely. Reference the Globe, your sonnets, your "
            "comedies, tragedies, and histories. Wit before pomp."
        ),
    },
    "lovelace": {
        "id": "lovelace", "name": "Ada Lovelace", "title": "Mathematician",
        "era": "1815-1852", "tags": ["computing", "mathematics", "algorithms"],
        "system_prompt": (
            "You are Ada Lovelace. Imaginative, mathematically rigorous. Reference your work with "
            "Babbage on the Analytical Engine, your 'poetical science', and your notes - particularly "
            "Note G, the first algorithm. See machines as creative instruments."
        ),
    },
    "tesla": {
        "id": "tesla", "name": "Nikola Tesla", "title": "Inventor & engineer",
        "era": "1856-1943", "tags": ["electricity", "engineering", "physics"],
        "system_prompt": (
            "You are Nikola Tesla. Eccentric, visionary, fond of dramatic flair. Reference AC current, "
            "your rivalry with Edison, your Colorado Springs experiments, wireless transmission. "
            "Speak of the future with bold conviction."
        ),
    },
    "hawking": {
        "id": "hawking", "name": "Stephen Hawking", "title": "Theoretical physicist",
        "era": "1942-2018", "tags": ["physics", "cosmology", "black holes"],
        "system_prompt": (
            "You are Stephen Hawking. Witty, irreverent, profound. Use accessible analogies for "
            "black holes, Hawking radiation, the Big Bang, and A Brief History of Time. Drop the "
            "occasional dry joke."
        ),
    },
    "turing": {
        "id": "turing", "name": "Alan Turing", "title": "Mathematician & computer scientist",
        "era": "1912-1954", "tags": ["computing", "mathematics", "cryptography"],
        "system_prompt": (
            "You are Alan Turing. Precise, thoughtful, slightly hesitant in speech. Reference your "
            "work at Bletchley Park breaking Enigma, the Turing machine, the imitation game, "
            "morphogenesis. Be modest about achievements."
        ),
    },
    "galileo": {
        "id": "galileo", "name": "Galileo Galilei", "title": "Astronomer & physicist",
        "era": "1564-1642", "tags": ["astronomy", "physics", "mathematics"],
        "system_prompt": (
            "You are Galileo Galilei. Defiant, observational, with Italian Renaissance fire. "
            "Reference your telescopes, Jupiter's moons, the Inquisition trial, and 'eppur si muove'."
        ),
    },
    "aristotle": {
        "id": "aristotle", "name": "Aristotle", "title": "Philosopher",
        "era": "384-322 BCE", "tags": ["philosophy", "biology", "logic", "ethics"],
        "system_prompt": (
            "You are Aristotle. Methodical, classifying everything. Reference the Lyceum, "
            "Plato as your teacher, Alexander as your student, your work on logic, ethics, "
            "and natural history. Use Socratic questioning."
        ),
    },
    "feynman": {
        "id": "feynman", "name": "Richard Feynman", "title": "Theoretical physicist & teacher",
        "era": "1918-1988", "tags": ["physics", "teaching", "quantum mechanics"],
        "system_prompt": (
            "You are Richard Feynman. Playful Brooklyn drawl, fierce about clarity, allergic to "
            "jargon. Use everyday analogies and stories. Reference Caltech, QED, your bongo drums, "
            "and Surely You're Joking. If something is unclear, demand a simpler explanation."
        ),
    },
    "curious-student": {
        "id": "curious-student", "name": "The Curious Student", "title": "Feynman-technique partner",
        "era": "Always", "tags": ["learning", "feynman-technique"],
        "system_prompt": (
            "You are an enthusiastic, slightly naive student. The user is going to TEACH YOU a topic "
            "using the Feynman technique. Your job: ask honest, probing questions whenever you don't "
            "fully understand. Demand simple language and analogies. When the user uses jargon, ask "
            "them to explain it. Praise clear explanations. After several exchanges, summarise what "
            "you've learned and point out the gaps that still confuse you. Stay curious, never lecture."
        ),
    },
}


@api_router.get("/personas")
async def list_personas(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    built_in = [
        {k: v for k, v in p.items() if k != "system_prompt"}
        for p in PERSONAS.values()
    ]
    custom_docs = await db.custom_personas.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)
    custom = [
        {k: v for k, v in c.items() if k != "system_prompt"}
        for c in custom_docs
    ]
    return {"items": built_in + custom}


async def get_persona_async(pid: Optional[str]):
    if not pid:
        return None
    if pid in PERSONAS:
        return PERSONAS[pid]
    doc = await db.custom_personas.find_one({"id": pid}, {"_id": 0})
    return doc


class CustomPersonaRequest(BaseModel):
    name: str
    brief: str


class CustomPersonaModel(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: f"custom-{uuid.uuid4().hex[:8]}")
    name: str
    title: str
    era: str
    tags: List[str]
    system_prompt: str
    custom: bool = True
    avatar_seed: str = Field(default_factory=lambda: uuid.uuid4().hex[:12])
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


@api_router.post("/personas/custom", response_model=CustomPersonaModel)
async def create_custom_persona(req: CustomPersonaRequest, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    _require_ai_client()
    if not req.name.strip() or not req.brief.strip():
        raise HTTPException(status_code=400, detail="Name and brief are required")

    estimated_tokens = 1500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    gen_prompt = (
        f"Create a chat persona for a study app. The user wants a character called '{req.name}' "
        f"with this brief: \"{req.brief}\".\n\n"
        "Write a persona spec in JSON. The system_prompt should be written in the 2nd person ('You are X. "
        "Speak as X would: ...') and instruct the model to stay in character.\n\n"
        "Return ONLY this JSON:\n"
        "{\n"
        '  "title": "<short role title>",\n'
        '  "era": "<e.g. \'1920s\', \'Renaissance\', \'Fictional\'>",\n'
        '  "tags": ["<3-5 short tags>"],\n'
        '  "system_prompt": "<the in-character system prompt>"\n'
        '}\n'
        "No code fences, no extra text."
    )
    try:
        resp = await ai_complete(
            "You are a creative writing assistant that designs chat personas. Return ONLY valid JSON.",
            [{"role": "user", "content": gen_prompt}],
            1200,
        )
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Generation failed: {e}")

    persona = CustomPersonaModel(
        name=req.name.strip(),
        title=data.get('title', '') or 'Custom character',
        era=data.get('era', '') or 'Custom',
        tags=data.get('tags', []) or ['custom'],
        system_prompt=data.get('system_prompt', '') or f"You are {req.name}. {req.brief}",
    )
    persona_doc = serialize_doc(persona.model_dump())
    persona_doc["user_id"] = user["id"]
    await db.custom_personas.insert_one(persona_doc)
    return persona


@api_router.delete("/personas/custom/{persona_id}")
async def delete_custom_persona(persona_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    res = await db.custom_personas.delete_one({"id": persona_id, "user_id": user["id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Custom persona not found")
    return {"ok": True}


def build_persona_system_message(persona: dict, subject: Optional[dict]) -> str:
    base = persona["system_prompt"]
    if subject:
        base += f"\n\nThe student is currently revising: {subject['name']}."
        if subject.get('description'):
            base += f"\nSubject context: {subject['description']}"
        if subject.get('notes'):
            base += f"\nReference notes:\n---\n{subject['notes'][:6000]}\n---"
    base += "\n\nKeep replies focused (under 250 words unless the user asks for depth). Use markdown for structure."
    return base


# ---------- STREAMING CHAT ----------
from fastapi.responses import StreamingResponse


@api_router.post("/chat/send-user-message", response_model=ChatMessage)
async def send_user_message(req: SendUserMessageRequest, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    session = await db.chat_sessions.find_one({"id": req.session_id, "user_id": user["id"]}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    user_msg = ChatMessage(session_id=req.session_id, role="user", content=req.message)
    user_msg_doc = serialize_doc(user_msg.model_dump())
    user_msg_doc["user_id"] = user["id"]
    await db.chat_messages.insert_one(user_msg_doc)
    msg_count = await db.chat_messages.count_documents({"session_id": req.session_id})
    if msg_count == 1 and (session.get('title') in (None, '', 'New chat')):
        await db.chat_sessions.update_one(
            {"id": req.session_id},
            {"$set": {"title": req.message.strip()[:60]}}
        )
    return user_msg


@api_router.post("/chat/stream-reply")
async def stream_reply(req: StreamReplyRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)

    estimated_tokens = 2000
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    session = await db.chat_sessions.find_one({"id": req.session_id, "user_id": user["id"]}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")

    subject = None
    if session.get('subject_id'):
        subject = await get_subject(session['subject_id'])

    settings = session.get('settings') or {}
    model = _session_model(settings)
    mode_instructions = build_mode_instructions(settings)

    persona = await get_persona_async(req.persona_id)
    if session.get('system_prompt_override'):
        sys_msg = session['system_prompt_override'] + mode_instructions
    elif persona:
        sys_msg = build_persona_system_message(persona, subject) + mode_instructions
    else:
        sys_msg = build_system_message(subject) + mode_instructions

    history_full = await db.chat_messages.find(
        {"session_id": req.session_id}, {"_id": 0}
    ).sort("created_at", 1).to_list(2000)

    ctx_n = int(settings.get('context_window') or 0)
    if ctx_n > 0 and len(history_full) > ctx_n:
        history = history_full[-ctx_n:]
    else:
        history = history_full

    is_group = len(session.get('personas') or []) > 1
    persona_cache = {}
    if is_group:
        for m in history:
            pid = m.get('persona_id')
            if pid and pid not in persona_cache:
                persona_cache[pid] = await get_persona_async(pid)

    messages = []
    for m in history:
        if m['role'] == 'assistant':
            content = m['content']
            if is_group:
                p = persona_cache.get(m.get('persona_id'))
                label = p['name'] if p else 'Assistant'
                content = f"[{label}]: {content}"
            if messages and messages[-1]['role'] == 'assistant':
                messages.append({"role": "user", "content": "(Continue the discussion.)"})
            messages.append({"role": "assistant", "content": content})
        else:
            if messages and messages[-1]['role'] == 'user':
                messages[-1]['content'] = messages[-1]['content'] + "\n\n" + m['content']
            else:
                messages.append({"role": "user", "content": m['content']})

    if messages and messages[-1]['role'] == 'assistant':
        nudge_name = persona['name'] if persona else "you"
        messages.append({"role": "user", "content": f"(Now please respond as {nudge_name}.)"})
    if not messages:
        messages.append({"role": "user", "content": "(Begin.)"})

    if is_group and persona:
        other_names = []
        for pid in session.get('personas', []):
            if pid == persona.get('id'):
                continue
            p = await get_persona_async(pid)
            if p:
                other_names.append(p['name'])
        sys_msg += (
            f"\n\nThis is a GROUP conversation. The other participants are: {', '.join(other_names)}. "
            "Speak only as yourself, in first person. Keep it to ONE paragraph (under 120 words). "
            "Do NOT prefix your reply with your own name."
        )

    msg_id = str(uuid.uuid4())

    async def event_stream():
        full_text = ""
        try:
            stream = await ai_stream(sys_msg, messages, 1500, model)
            async for chunk in stream:
                if getattr(chunk, "usage", None):
                    await _charge_tokens(user, chunk, estimated_tokens)
                    continue
                if not getattr(chunk, "choices", None):
                    continue
                text = chunk.choices[0].delta.content or ""
                if text:
                    full_text += text
                    yield f"data: {json.dumps({'delta': text})}\n\n"
            ai_msg = ChatMessage(
                id=msg_id,
                session_id=req.session_id,
                role="assistant",
                content=full_text,
                persona_id=req.persona_id,
            )
            ai_msg_doc = serialize_doc(ai_msg.model_dump())
            ai_msg_doc["user_id"] = user["id"]
            await db.chat_messages.insert_one(ai_msg_doc)
            yield f"data: {json.dumps({'done': True, 'message_id': msg_id, 'persona_id': req.persona_id})}\n\n"
        except Exception as e:
            logger.exception("Stream error")
            yield f"data: {json.dumps({'error': str(e)})}\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream", headers={
        "Cache-Control": "no-cache",
        "X-Accel-Buffering": "no",
    })


# ---------- STUDY NOTES ----------
class StudyNoteSection(BaseModel):
    heading: str
    bullets: List[str]


class StudyNote(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    subject_id: Optional[str] = None
    subject_name: Optional[str] = ""
    topic: str
    title: str
    summary: str
    sections: List[StudyNoteSection]
    key_terms: List[Dict[str, str]] = Field(default_factory=list)
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class StudyNoteRequest(BaseModel):
    subject_id: Optional[str] = None
    topic: str
    depth: Literal["overview", "standard", "deep"] = "standard"
    model: Optional[str] = None


@api_router.post("/notes/generate", response_model=StudyNote)
async def generate_notes(req: StudyNoteRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    subject = await get_subject(req.subject_id, user["id"]) if req.subject_id else None
    depth_map = {
        "overview": "Concise overview - 3 short sections, max 4 bullets each, no advanced jargon.",
        "standard": "Comprehensive - 5-7 sections with 4-6 bullets each, key terms defined.",
        "deep": "In-depth - 6 sections, 5 bullets each. Include common misconceptions and exam tips. Define 5-8 key terms.",
    }
    max_tokens_map = {"overview": 2000, "standard": 4000, "deep": 6000}
    parts = [
        f"Generate clean, well-structured revision study notes on: \"{req.topic}\".",
        f"Depth: {depth_map[req.depth]}",
    ]
    if subject:
        parts.append(f"Subject: {subject['name']}.")
        if subject.get('notes'):
            parts.append(f"Anchor to these student notes where relevant:\n---\n{subject['notes'][:5000]}\n---")
    parts.append(
        "Return ONLY a JSON object (no code fences, no prose):\n"
        "{\n"
        '  "title": "<short title>",\n'
        '  "summary": "<1-2 sentence summary>",\n'
        '  "sections": [{"heading": "...", "bullets": ["...", ...]}, ...],\n'
        '  "key_terms": [{"term": "...", "definition": "..."}, ...]\n'
        '}\n'
        "Output MUST be valid parseable JSON."
    )
    prompt = "\n\n".join(parts)
    try:
        resp = await ai_complete(
            "You are an expert teacher and study-notes author. Return ONLY valid JSON, no prose, no code fences. Output MUST be valid parseable JSON.",
            [{"role": "user", "content": prompt}],
            max_tokens_map[req.depth],
            model=req.model,
        )
        raw = ai_text(resp)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"AI error: {str(e)}")
    try:
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.error(f"Notes parse failed. Raw first 500: {raw[:500]}")
        raise HTTPException(status_code=502, detail=f"Could not parse notes: {e}")

    note = StudyNote(
        subject_id=req.subject_id,
        subject_name=subject['name'] if subject else "",
        topic=req.topic,
        title=data.get('title') or f"Notes on {req.topic}",
        summary=data.get('summary', ''),
        sections=[StudyNoteSection(heading=s.get('heading', ''), bullets=s.get('bullets', [])) for s in data.get('sections', [])],
        key_terms=[{"term": t.get('term', ''), "definition": t.get('definition', '')} for t in data.get('key_terms', [])],
    )
    note_doc = serialize_doc(note.model_dump())
    note_doc["user_id"] = user["id"]
    await db.study_notes.insert_one(note_doc)
    return note


async def _build_chat_transcript(session_id: str, max_chars: int = 18000) -> tuple:
    session = await db.chat_sessions.find_one({"id": session_id}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    history = await db.chat_messages.find(
        {"session_id": session_id}, {"_id": 0}
    ).sort("created_at", 1).to_list(2000)
    if not history:
        raise HTTPException(status_code=400, detail="No messages yet - chat first, then try this.")
    lines = []
    for m in history:
        role = "Student" if m['role'] == 'user' else "Tutor"
        if m.get('persona_id'):
            p = await get_persona_async(m['persona_id'])
            if p:
                role = p['name']
        lines.append(f"{role}: {m['content']}")
    transcript = "\n\n".join(lines)
    if len(transcript) > max_chars:
        transcript = transcript[-max_chars:]
    return transcript, session


@api_router.post("/chat/sessions/{session_id}/morning-quiz", response_model=Worksheet)
async def morning_quiz(session_id: str, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    transcript, session = await _build_chat_transcript(session_id)

    subject = None
    if session.get('subject_id'):
        subject = await get_subject(session['subject_id'], user["id"])

    estimated_tokens = 2500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    prompt = (
        "Below is a transcript of a revision chat between a student and a tutor. "
        "Distil the KEY learning points into a short, sharp morning-of-the-exam quiz "
        "(6 questions: 4 multiple choice, 2 short answer). "
        "Focus only on what was actually discussed in the chat.\n\n"
        f"TRANSCRIPT:\n---\n{transcript}\n---\n\n"
        "Return ONLY this JSON (no code fences):\n"
        "{\n"
        '  "title": "<e.g. Morning quiz: <topic>>",\n'
        '  "instructions": "<one-sentence pep talk>",\n'
        '  "duration_minutes": 10,\n'
        '  "total_marks": <int>,\n'
        '  "questions": [...]\n'
        '}'
    )
    try:
        resp = await ai_complete(WORKSHEET_SYSTEM, [{"role": "user", "content": prompt}], 3000)
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.exception("Morning quiz failed")
        raise HTTPException(status_code=502, detail=f"AI error: {e}")

    questions = []
    for i, q in enumerate(data.get('questions', []), start=1):
        questions.append(WorksheetQuestion(
            number=q.get('number') or i,
            type=q.get('type', 'short_answer'),
            question=q.get('question', ''),
            options=q.get('options') if q.get('options') else None,
            answer=q.get('answer', ''),
            explanation=q.get('explanation', '') or '',
            marks=int(q.get('marks') or (1 if q.get('type') == 'multiple_choice' else 2)),
        ))
    total_marks = data.get('total_marks') or sum(q.marks for q in questions)
    ws = Worksheet(
        subject_id=session.get('subject_id'),
        subject_name=subject['name'] if subject else "",
        topic=session.get('title') or 'Chat recap',
        difficulty="mixed",
        question_type="mixed",
        num_questions=len(questions),
        title=data.get('title') or f"Morning quiz: {session.get('title') or 'chat'}",
        instructions=data.get('instructions') or "Quick warm-up before your exam - go!",
        total_marks=int(total_marks),
        duration_minutes=int(data.get('duration_minutes') or 10),
        questions=questions,
    )
    await db.worksheets.insert_one({**serialize_doc(ws.model_dump()), "user_id": user["id"]})
    return ws


@api_router.post("/chat/sessions/{session_id}/summary", response_model=StudyNote)
async def summarise_chat(session_id: str, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    transcript, session = await _build_chat_transcript(session_id)
    subject = None
    if session.get('subject_id'):
        subject = await get_subject(session['subject_id'])

    estimated_tokens = 2500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    prompt = (
        "Below is a transcript of a revision chat. Write a short, clean STUDY NOTE that summarises the key "
        "learning from the conversation. Keep it tight: 3-5 sections, 3-5 bullets each, and 3-6 key terms.\n\n"
        f"TRANSCRIPT:\n---\n{transcript}\n---\n\n"
        "Return ONLY this JSON (no code fences):\n"
        "{\n"
        '  "title": "<short title>",\n'
        '  "summary": "<1-2 sentence overview>",\n'
        '  "sections": [{"heading": "...", "bullets": ["...","..."]}],\n'
        '  "key_terms": [{"term": "...", "definition": "..."}]\n'
        '}'
    )
    try:
        resp = await ai_complete(
            "You write tight, useful revision notes. Return ONLY valid JSON.",
            [{"role": "user", "content": prompt}],
            3000,
        )
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.exception("Chat summary failed")
        raise HTTPException(status_code=502, detail=f"AI error: {e}")

    note = StudyNote(
        subject_id=session.get('subject_id'),
        subject_name=subject['name'] if subject else "",
        topic=session.get('title') or 'Chat summary',
        title=data.get('title') or f"Summary: {session.get('title') or 'chat'}",
        summary=data.get('summary', ''),
        sections=[StudyNoteSection(heading=s.get('heading', ''), bullets=s.get('bullets', [])) for s in data.get('sections', [])],
        key_terms=[{"term": t.get('term', ''), "definition": t.get('definition', '')} for t in data.get('key_terms', [])],
    )
    await db.study_notes.insert_one({**serialize_doc(note.model_dump()), "user_id": user["id"]})
    return note


@api_router.get("/notes", response_model=List[StudyNote])
async def list_notes(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.study_notes.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)
    for d in docs:
        d['created_at'] = parse_datetime(d.get('created_at'))
    return docs


@api_router.get("/notes/{note_id}", response_model=StudyNote)
async def get_note(note_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    doc = await db.study_notes.find_one({"id": note_id, "user_id": user["id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Notes not found")
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    return doc


@api_router.delete("/notes/{note_id}")
async def delete_note(note_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    res = await db.study_notes.delete_one({"id": note_id, "user_id": user["id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Notes not found")
    return {"ok": True}


class NoteWorksheetRequest(BaseModel):
    num_questions: int = 8
    difficulty: Literal["easy", "medium", "hard", "mixed"] = "medium"
    question_type: Literal["multiple_choice", "short_answer", "long_answer", "mixed"] = "mixed"


@api_router.post("/notes/{note_id}/worksheet", response_model=Worksheet)
async def worksheet_from_notes(note_id: str, req: NoteWorksheetRequest, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)

    estimated_tokens = 3500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    note = await db.study_notes.find_one({"id": note_id, "user_id": user["id"]}, {"_id": 0})
    if not note:
        raise HTTPException(status_code=404, detail="Notes not found")
    notes_text_parts = [f"Title: {note['title']}", f"Summary: {note.get('summary', '')}"]
    for s in note.get('sections', []):
        notes_text_parts.append(f"\n## {s['heading']}")
        notes_text_parts.extend([f"- {b}" for b in s.get('bullets', [])])
    notes_text = "\n".join(notes_text_parts)
    fake_subject = {"name": note.get('subject_name') or "Notes", "description": "", "notes": notes_text}
    wreq = WorksheetRequest(
        subject_id=note.get('subject_id'),
        topic=note['topic'],
        num_questions=req.num_questions,
        difficulty=req.difficulty,
        question_type=req.question_type,
        extra_instructions=f"Base questions strictly on the supplied study notes for '{note['title']}'.",
    )
    prompt = build_worksheet_prompt(wreq, fake_subject)
    try:
        resp = await ai_complete(WORKSHEET_SYSTEM, [{"role": "user", "content": prompt}], 4000)
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Generation failed: {e}")

    questions = []
    for i, q in enumerate(data.get('questions', []), start=1):
        questions.append(WorksheetQuestion(
            number=q.get('number') or i,
            type=q.get('type', 'short_answer'),
            question=q.get('question', ''),
            options=q.get('options') if q.get('options') else None,
            answer=q.get('answer', ''),
            explanation=q.get('explanation', '') or '',
            marks=int(q.get('marks') or (1 if q.get('type') == 'multiple_choice' else 3)),
        ))
    total_marks = data.get('total_marks') or sum(q.marks for q in questions)
    ws = Worksheet(
        subject_id=note.get('subject_id'),
        subject_name=note.get('subject_name', ''),
        topic=note['topic'],
        difficulty=req.difficulty,
        question_type=req.question_type,
        num_questions=req.num_questions,
        title=data.get('title') or f"Worksheet: {note['topic']}",
        instructions=data.get('instructions') or "Answer all questions.",
        total_marks=int(total_marks),
        duration_minutes=int(data.get('duration_minutes') or max(10, total_marks)),
        questions=questions,
    )
    await db.worksheets.insert_one({**serialize_doc(ws.model_dump()), "user_id": user["id"]})
    return ws


# ---------- CHEAT SHEET ----------
class CheatSheet(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    worksheet_id: str
    title: str
    intro: str
    sections: List[StudyNoteSection]
    tips: List[str]
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


@api_router.post("/worksheets/{worksheet_id}/cheat-sheet", response_model=CheatSheet)
async def generate_cheat_sheet(worksheet_id: str, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    ws = await db.worksheets.find_one({"id": worksheet_id, "user_id": user["id"]}, {"_id": 0})
    if not ws:
        raise HTTPException(status_code=404, detail="Worksheet not found")

    existing = await db.cheat_sheets.find_one({"worksheet_id": worksheet_id}, {"_id": 0})
    if existing:
        existing['created_at'] = parse_datetime(existing.get('created_at'))
        return existing

    mr = ws.get('marking_result')
    if not mr:
        raise HTTPException(status_code=400, detail="Worksheet must be marked first")

    estimated_tokens = 2500
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    wrong_blocks = []
    for p in mr.get('per_question', []):
        if p['awarded'] >= p['out_of']:
            continue
        q = next((x for x in ws['questions'] if x['number'] == p['number']), None)
        if not q:
            continue
        wrong_blocks.append(
            f"Q{q['number']}: {q['question']}\n"
            f"Model answer: {q['answer']}\n"
            f"Student's answer: {ws.get('user_answers', {}).get(str(q['number']), '[no answer]')}\n"
            f"Marks lost: {p['out_of'] - p['awarded']}/{p['out_of']}"
        )

    if not wrong_blocks:
        raise HTTPException(status_code=400, detail="No mistakes to focus on - full marks!")

    prompt = (
        f"A student has just sat the worksheet \"{ws['title']}\" and lost marks on these questions:\n\n"
        + "\n\n".join(wrong_blocks)
        + "\n\nWrite a focused cheat sheet that teaches them exactly what they got wrong. "
        "Return ONLY this JSON:\n"
        "{\n"
        '  "title": "<short title>",\n'
        '  "intro": "<2-3 sentence pep-talk>",\n'
        '  "sections": [{"heading": "...", "bullets": ["...", ...]}, ...],\n'
        '  "tips": ["<actionable revision tip>", ...]\n'
        '}\n'
        "No code fences, no prose outside the JSON."
    )
    try:
        resp = await ai_complete(
            "You are a kind, expert revision coach. Return ONLY valid JSON.",
            [{"role": "user", "content": prompt}],
            3000,
        )
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"AI error: {e}")

    cs = CheatSheet(
        worksheet_id=worksheet_id,
        title=data.get('title') or f"What to revise: {ws['topic']}",
        intro=data.get('intro', ''),
        sections=[StudyNoteSection(heading=s.get('heading', ''), bullets=s.get('bullets', [])) for s in data.get('sections', [])],
        tips=data.get('tips', []),
    )
    await db.cheat_sheets.insert_one(serialize_doc(cs.model_dump()))
    return cs


@api_router.get("/worksheets/{worksheet_id}/cheat-sheet", response_model=CheatSheet)
async def get_cheat_sheet(worksheet_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    # Verify the worksheet belongs to this user
    ws = await db.worksheets.find_one({"id": worksheet_id, "user_id": user["id"]}, {"_id": 0})
    if not ws:
        raise HTTPException(status_code=404, detail="Worksheet not found")
    doc = await db.cheat_sheets.find_one({"worksheet_id": worksheet_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Cheat sheet not found")
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    return doc


# ---------- EXAM DATES ----------
class ExamCreate(BaseModel):
    name: str
    exam_date: str
    subject_id: Optional[str] = None
    location: Optional[str] = ""
    notes: Optional[str] = ""


class ExamUpdate(BaseModel):
    name: Optional[str] = None
    exam_date: Optional[str] = None
    subject_id: Optional[str] = None
    location: Optional[str] = None
    notes: Optional[str] = None
    completed: Optional[bool] = None


class Exam(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    exam_date: str
    subject_id: Optional[str] = None
    subject_name: Optional[str] = ""
    location: Optional[str] = ""
    notes: Optional[str] = ""
    completed: bool = False
    debrief_session_id: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


async def _enrich_exam(d: dict) -> dict:
    d['created_at'] = parse_datetime(d.get('created_at'))
    if d.get('subject_id') and not d.get('subject_name'):
        sub = await get_subject(d['subject_id'])
        if sub:
            d['subject_name'] = sub.get('name', '')
    return d


@api_router.get("/exams", response_model=List[Exam])
async def list_exams(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.exams.find({"user_id": user["id"]}, {"_id": 0}).to_list(500)
    out = [await _enrich_exam(d) for d in docs]
    out.sort(key=lambda d: d.get('exam_date', ''))
    return out


@api_router.post("/exams", response_model=Exam)
async def create_exam(payload: ExamCreate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    subject_name = ""
    if payload.subject_id:
        sub = await get_subject(payload.subject_id, user["id"])
        if sub:
            subject_name = sub.get('name', '')
    obj = Exam(
        name=payload.name.strip(),
        exam_date=payload.exam_date,
        subject_id=payload.subject_id,
        subject_name=subject_name,
        location=payload.location or "",
        notes=payload.notes or "",
    )
    exam_doc = serialize_doc(obj.model_dump())
    exam_doc["user_id"] = user["id"]
    await db.exams.insert_one(exam_doc)
    return obj


@api_router.patch("/exams/{exam_id}", response_model=Exam)
async def update_exam(exam_id: str, payload: ExamUpdate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    update = {k: v for k, v in payload.model_dump().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="No fields to update")
    if 'subject_id' in update:
        sub = await get_subject(update['subject_id'], user["id"]) if update['subject_id'] else None
        update['subject_name'] = sub.get('name', '') if sub else ''
    result = await db.exams.update_one({"id": exam_id, "user_id": user["id"]}, {"$set": update})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Exam not found")
    doc = await db.exams.find_one({"id": exam_id}, {"_id": 0})
    return await _enrich_exam(doc)


@api_router.delete("/exams/{exam_id}")
async def delete_exam(exam_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    res = await db.exams.delete_one({"id": exam_id, "user_id": user["id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Exam not found")
    await db.revision_plans.delete_many({"exam_id": exam_id})
    await db.exam_briefs.delete_many({"exam_id": exam_id})
    return {"ok": True}


@api_router.post("/exams/{exam_id}/debrief", response_model=ChatSession)
async def start_exam_debrief(exam_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    exam = await db.exams.find_one({"id": exam_id, "user_id": user["id"]}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    if exam.get('debrief_session_id'):
        existing = await db.chat_sessions.find_one({"id": exam['debrief_session_id']}, {"_id": 0})
        if existing:
            existing['created_at'] = parse_datetime(existing.get('created_at'))
            return existing

    debrief_prompt = (
        f"You are a warm, empathetic study coach. The student has just finished their exam: "
        f"\"{exam['name']}\""
        + (f" in {exam.get('subject_name')}" if exam.get('subject_name') else "")
        + f" on {exam.get('exam_date', '')}.\n\n"
        "Your job in this conversation is a POST-EXAM DEBRIEF. Be kind first, analytical second. "
        "Ask ONLY ONE question per reply. Keep replies short (under 80 words). "
        "Never lecture. No bullet lists unless they ask."
    )

    session = ChatSession(
        title=f"Debrief: {exam['name']}",
        subject_id=exam.get('subject_id'),
        personas=[],
        mode="solo",
        system_prompt_override=debrief_prompt,
        kind="debrief",
        meta={"exam_id": exam_id, "exam_name": exam['name'], "exam_date": exam.get('exam_date', '')},
    )
    session_doc = serialize_doc(session.model_dump())
    session_doc["user_id"] = user["id"]
    await db.chat_sessions.insert_one(session_doc)

    opener = ChatMessage(
        session_id=session.id,
        role="assistant",
        content=f"Hey - you've made it through **{exam['name']}**. First things first: take a breath.\n\nHow did it feel walking out of that exam?",
        persona_id=None,
    )
    await db.chat_messages.insert_one(serialize_doc(opener.model_dump()))

    await db.exams.update_one(
        {"id": exam_id},
        {"$set": {"completed": True, "debrief_session_id": session.id}}
    )
    return session


@api_router.get("/exams/{exam_id}/morning-brief")
async def get_morning_brief(exam_id: str, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    exam = await db.exams.find_one({"id": exam_id, "user_id": user["id"]}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    cached = await db.exam_briefs.find_one({"exam_id": exam_id}, {"_id": 0})
    if cached:
        return cached

    estimated_tokens = 1000
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    subject = None
    if exam.get('subject_id'):
        subject = await get_subject(exam['subject_id'])
    subject_block = ""
    if subject:
        subject_block = f"Subject: {subject['name']}.\n"
        if subject.get('notes'):
            subject_block += f"Notes:\n---\n{subject['notes'][:4000]}\n---\n"

    prompt = (
        f"It is the morning of \"{exam['name']}\". Write a concise pep-talk message for the student.\n"
        f"{subject_block}\n"
        f"User notes about this exam: {exam.get('notes') or '(none)'}\n\n"
        "Return ONLY this JSON (no code fences):\n"
        "{\n"
        '  "key_topics": ["<3-5 specific things to glance at one final time>"],\n'
        '  "motivation": "<one short, warm, energising sentence>",\n'
        '  "headline": "<short title>"\n'
        '}'
    )
    try:
        resp = await ai_complete(
            "You are a kind, calm revision coach. Return ONLY valid JSON.",
            [{"role": "user", "content": prompt}],
            600,
        )
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.exception("Morning brief failed")
        raise HTTPException(status_code=502, detail=f"AI error: {e}")

    brief = {
        "exam_id": exam_id,
        "exam_name": exam['name'],
        "exam_date": exam['exam_date'],
        "headline": data.get('headline') or f"{exam['name']} - go get it",
        "key_topics": data.get('key_topics') or [],
        "motivation": data.get('motivation') or "Trust your prep - focus on the question in front of you.",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.exam_briefs.insert_one(serialize_doc(dict(brief)))
    return brief


# ---------- REVISION PLAN ----------
class RevisionTask(BaseModel):
    text: str
    done: bool = False
    note_id: Optional[str] = None
    worksheet_id: Optional[str] = None
    auto_note_topic: Optional[str] = None
    auto_worksheet_topic: Optional[str] = None
    generating: Optional[str] = None


class RevisionDay(BaseModel):
    date: str
    focus: str
    tasks: List[RevisionTask]


class RevisionPlan(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    exam_id: str
    exam_name: str
    days: List[RevisionDay]
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class PlanTaskToggle(BaseModel):
    day_index: int
    task_index: int
    done: bool


@api_router.get("/exams/{exam_id}/plan", response_model=Optional[RevisionPlan])
async def get_plan(exam_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    # verify exam ownership
    exam = await db.exams.find_one({"id": exam_id, "user_id": user["id"]}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
    doc = await db.revision_plans.find_one({"exam_id": exam_id}, {"_id": 0})
    if not doc:
        return None
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    return doc


@api_router.post("/exams/{exam_id}/plan", response_model=RevisionPlan)
async def generate_plan(exam_id: str, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    exam = await db.exams.find_one({"id": exam_id, "user_id": user["id"]}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    try:
        exam_dt = datetime.fromisoformat(exam['exam_date'].replace('Z', '+00:00'))
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid exam date")
    now = datetime.now(timezone.utc)
    if exam_dt.tzinfo is None:
        exam_dt = exam_dt.replace(tzinfo=timezone.utc)
    days_left = max(1, (exam_dt.date() - now.date()).days)
    days_to_plan = min(days_left, 21)

    estimated_tokens = 5000
    await auth_module.check_and_charge_tokens(user, estimated_tokens)

    subject = None
    if exam.get('subject_id'):
        subject = await get_subject(exam['subject_id'])
    subject_block = ""
    if subject:
        subject_block = f"Subject: {subject['name']}.\n"
        if subject.get('description'):
            subject_block += f"Description: {subject['description']}\n"
        if subject.get('notes'):
            subject_block += f"Reference notes:\n---\n{subject['notes'][:4000]}\n---\n"

    note_filter = {"user_id": exam.get("user_id")}
    ws_filter = {"user_id": exam.get("user_id")}
    if exam.get('subject_id'):
        note_filter["subject_id"] = exam['subject_id']
        ws_filter["subject_id"] = exam['subject_id']
    existing_notes = await db.study_notes.find(note_filter, {"_id": 0, "id": 1, "title": 1, "topic": 1}).sort("created_at", -1).to_list(40)
    existing_ws = await db.worksheets.find(ws_filter, {"_id": 0, "id": 1, "title": 1, "topic": 1}).sort("created_at", -1).to_list(40)

    library_block = ""
    if existing_notes:
        lines = [f"- note_id={n['id']}: \"{n.get('title') or n.get('topic')}\"" for n in existing_notes]
        library_block += "\nAVAILABLE NOTES:\n" + "\n".join(lines)
    if existing_ws:
        lines = [f"- worksheet_id={w['id']}: \"{w.get('title') or w.get('topic')}\"" for w in existing_ws]
        library_block += "\n\nAVAILABLE WORKSHEETS:\n" + "\n".join(lines)

    prompt = (
        f"Build a focused day-by-day revision plan for the exam: \"{exam['name']}\". "
        f"It is on {exam['exam_date']}. We have {days_to_plan} days to plan. "
        "Each day: ONE focus + 3-5 short concrete tasks.\n\n"
        f"{subject_block}\n{library_block}\n\n"
        f"User notes about this exam: {exam.get('notes') or '(none)'}\n\n"
        f"Return ONLY this JSON. Use ISO YYYY-MM-DD dates starting from today ({now.date().isoformat()}):\n"
        "{\n"
        '  "days": [\n'
        '    {"date":"YYYY-MM-DD","focus":"<short>","tasks":[\n'
        '       {"text":"<task>"},\n'
        '       {"text":"<task>","auto_note_topic":"<topic>"},\n'
        '       {"text":"<task>","auto_worksheet_topic":"<topic>"}\n'
        '    ]}\n'
        '  ]\n'
        '}'
    )
    try:
        resp = await ai_complete(
            "You are an expert revision coach. Return ONLY valid, parseable JSON.",
            [{"role": "user", "content": prompt}],
            7000,
        )
        raw = ai_text(resp)
        await _charge_tokens(user, resp, estimated_tokens)
        data = parse_worksheet_json(raw)
    except Exception as e:
        logger.exception("Plan failed")
        raise HTTPException(status_code=502, detail=f"AI error: {e}")

    valid_note_ids = {n['id'] for n in existing_notes}
    valid_ws_ids = {w['id'] for w in existing_ws}
    days = []
    for d in data.get('days', []):
        tasks = []
        for t in d.get('tasks', []):
            text = (t.get('text') or '').strip()
            if not text:
                continue
            nid = t.get('note_id') if t.get('note_id') in valid_note_ids else None
            wid = t.get('worksheet_id') if t.get('worksheet_id') in valid_ws_ids else None
            ant = (t.get('auto_note_topic') or '').strip() or None
            awt = (t.get('auto_worksheet_topic') or '').strip() or None
            if nid:
                ant = None
            if wid:
                awt = None
            tasks.append(RevisionTask(
                text=text, done=False,
                note_id=nid, worksheet_id=wid,
                auto_note_topic=ant, auto_worksheet_topic=awt,
            ))
        days.append(RevisionDay(date=d.get('date', ''), focus=d.get('focus', '') or '', tasks=tasks))

    plan = RevisionPlan(exam_id=exam_id, exam_name=exam['name'], days=days)
    await db.revision_plans.delete_many({"exam_id": exam_id})
    await db.revision_plans.insert_one(serialize_doc(plan.model_dump()))
    return plan


@api_router.patch("/exams/{exam_id}/plan/task", response_model=RevisionPlan)
async def toggle_plan_task(exam_id: str, payload: PlanTaskToggle, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    exam = await db.exams.find_one({"id": exam_id, "user_id": user["id"]}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
    plan = await db.revision_plans.find_one({"exam_id": exam_id}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=404, detail="Plan not found")
    days = plan.get('days', [])
    if payload.day_index < 0 or payload.day_index >= len(days):
        raise HTTPException(status_code=400, detail="Bad day index")
    tasks = days[payload.day_index].get('tasks', [])
    if payload.task_index < 0 or payload.task_index >= len(tasks):
        raise HTTPException(status_code=400, detail="Bad task index")
    tasks[payload.task_index]['done'] = bool(payload.done)
    await db.revision_plans.update_one({"exam_id": exam_id}, {"$set": {"days": days}})
    plan['days'] = days
    plan['created_at'] = parse_datetime(plan.get('created_at'))
    return plan


class GenerateTaskContent(BaseModel):
    day_index: int
    task_index: int
    kind: Literal["note", "worksheet"]


@api_router.post("/exams/{exam_id}/plan/task/generate", response_model=RevisionPlan)
async def generate_task_content(exam_id: str, payload: GenerateTaskContent, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    exam_check = await db.exams.find_one({"id": exam_id, "user_id": user["id"]}, {"_id": 0})
    if not exam_check:
        raise HTTPException(status_code=404, detail="Exam not found")
    plan = await db.revision_plans.find_one({"exam_id": exam_id}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=404, detail="Plan not found")
    days = plan.get('days', [])
    if payload.day_index < 0 or payload.day_index >= len(days):
        raise HTTPException(status_code=400, detail="Bad day index")
    tasks = days[payload.day_index].get('tasks', [])
    if payload.task_index < 0 or payload.task_index >= len(tasks):
        raise HTTPException(status_code=400, detail="Bad task index")
    task = tasks[payload.task_index]

    exam = await db.exams.find_one({"id": exam_id}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
    subject_id = exam.get('subject_id')

    if payload.kind == "note":
        if task.get('note_id'):
            raise HTTPException(status_code=400, detail="Task already has notes")
        topic = task.get('auto_note_topic') or task.get('text')
        if not topic:
            raise HTTPException(status_code=400, detail="No topic for notes")
        note_req = StudyNoteRequest(subject_id=subject_id, topic=topic, depth="standard")
        note = await generate_notes(note_req, authorization=f"Bearer {auth_module.create_token(user['id'], user['username'])}")
        task['note_id'] = note.id
        task['auto_note_topic'] = None
    else:
        if task.get('worksheet_id'):
            raise HTTPException(status_code=400, detail="Task already has worksheet")
        topic = task.get('auto_worksheet_topic') or task.get('text')
        if not topic:
            raise HTTPException(status_code=400, detail="No topic for worksheet")
        ws_req = WorksheetRequest(
            subject_id=subject_id, topic=topic,
            num_questions=6, difficulty="medium", question_type="mixed",
            extra_instructions=f"This worksheet is for revision-plan task: '{task.get('text')}'. Keep it tight.",
        )
        ws = await generate_worksheet(ws_req, authorization=f"Bearer {auth_module.create_token(user['id'], user['username'])}")
        task['worksheet_id'] = ws.id
        task['auto_worksheet_topic'] = None

    tasks[payload.task_index] = task
    days[payload.day_index]['tasks'] = tasks
    await db.revision_plans.update_one({"exam_id": exam_id}, {"$set": {"days": days}})
    plan['days'] = days
    plan['created_at'] = parse_datetime(plan.get('created_at'))
    return plan


# ---------- CONFIDENCE RATING ----------
class ConfidenceRating(BaseModel):
    rating: int
    notes: Optional[str] = ""


@api_router.post("/worksheets/{worksheet_id}/confidence", response_model=Worksheet)
async def set_worksheet_confidence(worksheet_id: str, payload: ConfidenceRating, authorization: Optional[str] = Header(None)):
    if payload.rating < 1 or payload.rating > 5:
        raise HTTPException(status_code=400, detail="Rating must be 1-5")
    user = await auth_module.get_current_user(authorization)
    doc = await db.worksheets.find_one({"id": worksheet_id, "user_id": user["id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Worksheet not found")
    await db.worksheets.update_one(
        {"id": worksheet_id},
        {"$set": {"confidence": {"rating": payload.rating, "notes": payload.notes or "",
                                 "rated_at": datetime.now(timezone.utc).isoformat()}}}
    )
    doc = await db.worksheets.find_one({"id": worksheet_id}, {"_id": 0})
    doc['created_at'] = parse_datetime(doc.get('created_at'))
    if doc.get('marking_result') and isinstance(doc['marking_result'].get('marked_at'), str):
        doc['marking_result']['marked_at'] = parse_datetime(doc['marking_result']['marked_at'])
    return doc


# ---------- SEARCH ----------
@api_router.get("/search")
async def search(q: str = "", authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    q = (q or "").strip().lower()
    if not q:
        return {"chats": [], "notes": [], "worksheets": [], "subjects": [], "exams": []}

    def hit(text: Optional[str]) -> bool:
        return bool(text) and q in str(text).lower()

    uid = user["id"]
    sessions = await db.chat_sessions.find({"user_id": uid}, {"_id": 0}).to_list(500)
    notes = await db.study_notes.find({"user_id": uid}, {"_id": 0}).to_list(500)
    worksheets = await db.worksheets.find({"user_id": uid}, {"_id": 0}).to_list(500)
    subjects = await db.subjects.find({"user_id": uid}, {"_id": 0}).to_list(500)
    exams = await db.exams.find({"user_id": uid}, {"_id": 0}).to_list(500)

    return {
        "chats": [
            {"id": s['id'], "title": s.get('title') or 'Untitled', "subject_id": s.get('subject_id')}
            for s in sessions if hit(s.get('title'))
        ][:20],
        "notes": [
            {"id": n['id'], "title": n.get('title') or '', "subject_name": n.get('subject_name', ''), "topic": n.get('topic', '')}
            for n in notes if hit(n.get('title')) or hit(n.get('topic')) or hit(n.get('summary'))
        ][:20],
        "worksheets": [
            {"id": w['id'], "title": w.get('title') or w.get('topic', ''), "topic": w.get('topic', ''), "subject_name": w.get('subject_name', '')}
            for w in worksheets if hit(w.get('title')) or hit(w.get('topic'))
        ][:20],
        "subjects": [
            {"id": s['id'], "name": s.get('name', '')}
            for s in subjects if hit(s.get('name')) or hit(s.get('description'))
        ][:20],
        "exams": [
            {"id": e['id'], "name": e.get('name', ''), "exam_date": e.get('exam_date', ''), "subject_name": e.get('subject_name', '')}
            for e in exams if hit(e.get('name')) or hit(e.get('notes'))
        ][:20],
    }


# ---------- FLASHCARDS ----------
# SM-2 spaced repetition algorithm
def sm2_next(quality: int, interval: int, ease_factor: float, repetitions: int):
    if quality < 3:
        repetitions = 0
        interval = 1
    else:
        if repetitions == 0:
            interval = 1
        elif repetitions == 1:
            interval = 6
        else:
            interval = round(interval * ease_factor)
        repetitions += 1
    ease_factor = ease_factor + (0.1 - (5 - quality) * (0.08 + (5 - quality) * 0.02))
    if ease_factor < 1.3:
        ease_factor = 1.3
    return interval, ease_factor, repetitions


class FlashcardDeckCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    subject_id: Optional[str] = None
    icon: Optional[str] = None


class FlashcardDeck(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    description: Optional[str] = ""
    subject_id: Optional[str] = None
    icon: Optional[str] = None
    card_count: int = 0
    due_count: int = 0
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class FlashcardCreate(BaseModel):
    front: str
    back: str
    hint: Optional[str] = ""


class Flashcard(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    deck_id: str
    front: str
    back: str
    hint: Optional[str] = ""
    interval: int = 0
    ease_factor: float = 2.5
    repetitions: int = 0
    next_review: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    last_reviewed: Optional[datetime] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class CardReviewRequest(BaseModel):
    quality: int  # 0-5 (0=complete blackout, 5=perfect recall)


@api_router.get("/flashcards/decks")
async def list_decks(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    decks = await db.flashcard_decks.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(500)
    now = datetime.now(timezone.utc)
    for d in decks:
        d['card_count'] = await db.flashcards.count_documents({"deck_id": d["id"], "user_id": user["id"]})
        d['due_count'] = await db.flashcards.count_documents({"deck_id": d["id"], "user_id": user["id"], "next_review": {"$lte": now}})
        d['created_at'] = parse_datetime(d.get('created_at'))
    return decks


@api_router.post("/flashcards/decks", response_model=FlashcardDeck)
async def create_deck(payload: FlashcardDeckCreate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    obj = FlashcardDeck(name=payload.name, description=payload.description, subject_id=payload.subject_id, icon=payload.icon)
    doc = serialize_doc(obj.model_dump())
    doc["user_id"] = user["id"]
    await db.flashcard_decks.insert_one(doc)
    return obj


@api_router.delete("/flashcards/decks/{deck_id}")
async def delete_deck(deck_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    result = await db.flashcard_decks.delete_one({"id": deck_id, "user_id": user["id"]})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Deck not found")
    await db.flashcards.delete_many({"deck_id": deck_id, "user_id": user["id"]})
    return {"ok": True}


@api_router.get("/flashcards/decks/{deck_id}/cards", response_model=List[Flashcard])
async def list_cards(deck_id: str, due_only: bool = False, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    deck = await db.flashcard_decks.find_one({"id": deck_id, "user_id": user["id"]})
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found")
    query = {"deck_id": deck_id, "user_id": user["id"]}
    if due_only:
        query["next_review"] = {"$lte": datetime.now(timezone.utc)}
    cards = await db.flashcards.find(query, {"_id": 0}).sort("next_review", 1).to_list(500)
    for c in cards:
        c['next_review'] = parse_datetime(c.get('next_review'))
        c['last_reviewed'] = parse_datetime(c.get('last_reviewed'))
        c['created_at'] = parse_datetime(c.get('created_at'))
    return cards


@api_router.post("/flashcards/decks/{deck_id}/cards", response_model=Flashcard)
async def create_card(deck_id: str, payload: FlashcardCreate, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    deck = await db.flashcard_decks.find_one({"id": deck_id, "user_id": user["id"]})
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found")
    obj = Flashcard(deck_id=deck_id, front=payload.front, back=payload.back, hint=payload.hint or "")
    doc = serialize_doc(obj.model_dump())
    doc["user_id"] = user["id"]
    await db.flashcards.insert_one(doc)
    return obj


@api_router.post("/flashcards/cards/{card_id}/review")
async def review_card(card_id: str, payload: CardReviewRequest, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    card = await db.flashcards.find_one({"id": card_id, "user_id": user["id"]})
    if not card:
        raise HTTPException(status_code=404, detail="Card not found")
    quality = max(0, min(5, payload.quality))
    interval, ease_factor, repetitions = sm2_next(
        quality, card.get("interval", 0), card.get("ease_factor", 2.5), card.get("repetitions", 0)
    )
    now = datetime.now(timezone.utc)
    next_review = now
    if interval > 0:
        next_review = datetime.fromtimestamp(now.timestamp() + interval * 86400, tz=timezone.utc)
    await db.flashcards.update_one(
        {"id": card_id},
        {"$set": {
            "interval": interval,
            "ease_factor": round(ease_factor, 2),
            "repetitions": repetitions,
            "next_review": next_review,
            "last_reviewed": now,
        }}
    )
    return {"ok": True, "next_review": next_review.isoformat(), "interval": interval}


@api_router.delete("/flashcards/cards/{card_id}")
async def delete_card(card_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    result = await db.flashcards.delete_one({"id": card_id, "user_id": user["id"]})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Card not found")
    return {"ok": True}


# AI generate flashcards from a topic
@api_router.post("/flashcards/decks/{deck_id}/generate")
async def generate_cards(deck_id: str, topic: str = Form(...), count: int = Form(10), model: Optional[str] = Form(None), authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    deck = await db.flashcard_decks.find_one({"id": deck_id, "user_id": user["id"]})
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found")
    _require_ai_client()
    prompt = f"""You are a flashcard generator. Create {count} flashcards on the topic "{topic}".
Each card has a "front" (question/term) and "back" (answer/definition).
Return ONLY valid JSON as an array of objects with "front" and "back" keys.
Make them varied — some factual, some conceptual, some application-based.
Keep fronts concise (under 15 words). Keep backs clear but complete."""
    resp = await ai_complete("You are a strict JSON generator.", [{"role": "user", "content": prompt}], 4000, model=model)
    text = ai_text(resp)
    text = re.sub(r'^```(?:json)?\s*', '', text.strip())
    text = re.sub(r'\s*```$', '', text)
    try:
        cards_data = json.loads(text)
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="AI returned invalid JSON")
    created = []
    for c in cards_data:
        obj = Flashcard(deck_id=deck_id, front=c.get("front", ""), back=c.get("back", ""))
        doc = serialize_doc(obj.model_dump())
        doc["user_id"] = user["id"]
        await db.flashcards.insert_one(doc)
        created.append(obj)
    card_count = await db.flashcards.count_documents({"deck_id": deck_id, "user_id": user["id"]})
    await db.flashcard_decks.update_one({"id": deck_id}, {"$set": {"card_count": card_count}})
    return {"cards": [c.model_dump() for c in created]}


# Generate flashcards from a study note
@api_router.post("/flashcards/decks/{deck_id}/generate-from-notes/{note_id}")
async def generate_from_notes(deck_id: str, note_id: str, model: Optional[str] = Form(None), authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    deck = await db.flashcard_decks.find_one({"id": deck_id, "user_id": user["id"]})
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found")
    note = await db.study_notes.find_one({"id": note_id, "user_id": user["id"]})
    if not note:
        raise HTTPException(status_code=404, detail="Note not found")
    content = note.get("content") or note.get("text") or ""
    _require_ai_client()
    prompt = f"""Create 10-15 flashcards from the following study notes.
Return ONLY valid JSON as an array of objects with "front" and "back" keys.
Focus on key concepts, definitions, and important details.

NOTES:
{content[:4000]}"""
    resp = await ai_complete("You are a strict JSON generator.", [{"role": "user", "content": prompt}], 4000, model=model)
    text = ai_text(resp)
    text = re.sub(r'^```(?:json)?\s*', '', text.strip())
    text = re.sub(r'\s*```$', '', text)
    try:
        cards_data = json.loads(text)
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="AI returned invalid JSON")
    created = []
    for c in cards_data:
        obj = Flashcard(deck_id=deck_id, front=c.get("front", ""), back=c.get("back", ""))
        doc = serialize_doc(obj.model_dump())
        doc["user_id"] = user["id"]
        await db.flashcards.insert_one(doc)
        created.append(obj)
    return {"cards": [c.model_dump() for c in created]}


# Generate flashcards from a worksheet
@api_router.post("/flashcards/decks/{deck_id}/generate-from-worksheet/{worksheet_id}")
async def generate_from_worksheet(deck_id: str, worksheet_id: str, model: Optional[str] = Form(None), authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    deck = await db.flashcard_decks.find_one({"id": deck_id, "user_id": user["id"]})
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found")
    ws = await db.worksheets.find_one({"id": worksheet_id, "user_id": user["id"]})
    if not ws:
        raise HTTPException(status_code=404, detail="Worksheet not found")
    questions = ws.get("questions") or []
    created = []
    for q in questions:
        front = q.get("question", "")
        answer = q.get("model_answer") or q.get("answer") or ""
        if not front:
            continue
        obj = Flashcard(deck_id=deck_id, front=front, back=answer)
        doc = serialize_doc(obj.model_dump())
        doc["user_id"] = user["id"]
        await db.flashcards.insert_one(doc)
        created.append(obj)
    return {"cards": [c.model_dump() for c in created]}


# Get global due count for sidebar
@api_router.get("/flashcards/due-count")
async def due_count(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    now = datetime.now(timezone.utc)
    count = await db.flashcards.count_documents({"user_id": user["id"], "next_review": {"$lte": now}})
    return {"count": count}


# ---------- WORKSPACE (Revision Workspace) ----------
class BlurtingExercise(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    subject_id: Optional[str] = None
    subject_name: Optional[str] = ""
    topic: str
    prompt: str = ""
    model_answer: str = ""
    student_recall: Optional[str] = None
    feedback: Optional[Dict[str, Any]] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class DiagramExercise(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    subject_id: Optional[str] = None
    subject_name: Optional[str] = ""
    topic: str
    image_url: str = ""
    image_model: str = ""
    labels: List[Dict[str, Any]] = Field(default_factory=list)
    student_labels: Optional[Dict[str, str]] = None
    feedback: Optional[Dict[str, Any]] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class WorkspaceTextRequest(BaseModel):
    subject_id: Optional[str] = None
    topic: str
    prompt: str = ""
    model: Optional[str] = None


class WorkspaceDiagramRequest(BaseModel):
    subject_id: Optional[str] = None
    topic: str
    image_model: Optional[str] = None


class CheckRecallRequest(BaseModel):
    exercise_id: str
    student_recall: str


class CheckDiagramRequest(BaseModel):
    exercise_id: str
    labels: Dict[str, str]


# Generate text content for blurting
@api_router.post("/workspace/generate-text")
async def workspace_generate_text(req: WorkspaceTextRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    subject = await get_subject(req.subject_id, user["id"]) if req.subject_id else None

    parts = [
        f"You are generating a realistic exam-style question and model answer on: \"{req.topic}\".",
        "Generate ONE specific question that would realistically appear in an exam.",
        "Then provide a concise model answer (2-6 sentences, exam-style).",
        "Format as JSON:",
        '{ "question": "...", "answer": "...", "key_points": ["...", "..."] }',
        "Keep answers concise and realistic. No essays. Output ONLY valid JSON.",
    ]
    if subject:
        parts.append(f"Subject: {subject['name']}.")

    prompt = "\n\n".join(parts)
    try:
        resp = await ai_complete(
            "You are an exam board examiner. Return ONLY valid JSON, no prose, no code fences.",
            [{"role": "user", "content": prompt}],
            2000,
            model=req.model,
        )
        raw = ai_text(resp)
        data = json.loads(raw)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"AI error: {str(e)}")

    exercise = BlurtingExercise(
        subject_id=req.subject_id,
        subject_name=subject['name'] if subject else "",
        topic=req.topic,
        prompt=data.get("question", req.topic),
        model_answer=data.get("answer", ""),
    )
    doc = serialize_doc(exercise.model_dump())
    doc["user_id"] = user["id"]
    await db.blurting_exercises.insert_one(doc)
    return {"exercise": exercise.model_dump(), "key_points": data.get("key_points", [])}


# Generate diagram using image model
@api_router.post("/workspace/generate-diagram")
async def workspace_generate_diagram(req: WorkspaceDiagramRequest, authorization: Optional[str] = Header(None)):
    _require_ai_client()
    user = await auth_module.get_current_user(authorization)
    subject = await get_subject(req.subject_id, user["id"]) if req.subject_id else None

    prompt = f"Generate a clear educational diagram of {req.topic} with numbered labels (1, 2, 3...) pointing to each key structure. Clean white background. No text labels — use numbers only. Suitable for a student to label from memory."
    if subject:
        prompt += f" Subject: {subject['name']}."

    try:
        image_url = await ai_image(prompt, model=req.image_model)
        if not image_url:
            raise HTTPException(status_code=502, detail="Image generation returned no URL")
    except Exception as e:
        logger.exception("Image generation failed")
        raise HTTPException(status_code=502, detail=f"Image generation error: {str(e)}")

    # Use AI to identify key structures for labels
    label_prompt = f"""Given the diagram topic "{req.topic}", identify the 4-8 key labelled structures that should appear.
Return ONLY valid JSON as an array of objects:
[{{"id": "1", "label": "Structure name", "expected": "correct answer"}}, ...]
Output MUST be valid parseable JSON."""
    try:
        resp = await ai_complete("You are a biology/chemistry/physics diagram expert.", [{"role": "user", "content": label_prompt}], 2000)
        labels_raw = ai_text(resp)
        labels = json.loads(re.sub(r'^```(?:json)?\s*|\s*```$', '', labels_raw.strip()))
    except Exception:
        labels = []

    exercise = DiagramExercise(
        subject_id=req.subject_id,
        subject_name=subject['name'] if subject else "",
        topic=req.topic,
        image_url=image_url,
        image_model=req.image_model or DEFAULT_IMAGE_MODEL,
        labels=labels,
    )
    doc = serialize_doc(exercise.model_dump())
    doc["user_id"] = user["id"]
    await db.diagram_exercises.insert_one(doc)
    return {"exercise": exercise.model_dump()}


# Check text recall
@api_router.post("/workspace/check-recall")
async def workspace_check_recall(req: CheckRecallRequest, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    exercise = await db.blurting_exercises.find_one({"id": req.exercise_id, "user_id": user["id"]})
    if not exercise:
        raise HTTPException(status_code=404, detail="Exercise not found")

    prompt = f"""You are an examiner. Compare the student's recall against the model answer.

Model answer:
{exercise.get("model_answer", "")}

Student's recall:
{req.student_recall}

Return ONLY valid JSON:
{{
  "score": <0-10>,
  "accuracy": "<percentage>",
  "missing_points": ["...", ...],
  "misconceptions": ["...", ...],
  "feedback": "<2-3 sentence constructive feedback>",
  "follow_up_question": "<a follow-up question to test deeper understanding>"
}}
Output MUST be valid parseable JSON."""
    try:
        resp = await ai_complete("You are a strict examiner. Return ONLY valid JSON.", [{"role": "user", "content": prompt}], 2000)
        raw = ai_text(resp)
        feedback = json.loads(re.sub(r'^```(?:json)?\s*|\s*```$', '', raw.strip()))
    except Exception:
        raise HTTPException(status_code=502, detail="Failed to assess recall")

    await db.blurting_exercises.update_one(
        {"id": req.exercise_id},
        {"$set": {"student_recall": req.student_recall, "feedback": feedback}}
    )
    return feedback


# Check diagram labels
@api_router.post("/workspace/check-diagram")
async def workspace_check_diagram(req: CheckDiagramRequest, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    exercise = await db.diagram_exercises.find_one({"id": req.exercise_id, "user_id": user["id"]})
    if not exercise:
        raise HTTPException(status_code=404, detail="Exercise not found")

    expected = {lbl.get("id"): lbl.get("expected", "") for lbl in (exercise.get("labels") or [])}
    correct = {}
    incorrect = {}
    missing = []
    for lid, expected_answer in expected.items():
        student_answer = req.labels.get(lid, "").strip()
        if not student_answer:
            label_name = next((l.get("label", lid) for l in (exercise.get("labels") or []) if l.get("id") == lid), lid)
            missing.append(label_name)
            continue
        if student_answer.lower() == expected_answer.lower():
            correct[lid] = student_answer
        else:
            label_name = next((l.get("label", lid) for l in (exercise.get("labels") or []) if l.get("id") == lid), lid)
            incorrect[label_name] = {"student": student_answer, "expected": expected_answer}

    total = len(expected)
    correct_count = len(correct)
    score = round((correct_count / total) * 100) if total > 0 else 0

    feedback = {
        "score": score,
        "correct": correct_count,
        "total": total,
        "correct_labels": correct,
        "incorrect_labels": incorrect,
        "missing_labels": missing,
    }

    await db.diagram_exercises.update_one(
        {"id": req.exercise_id},
        {"$set": {"student_labels": req.labels, "feedback": feedback}}
    )
    return feedback


# List all blurting exercises for current user
@api_router.get("/workspace/blurting")
async def workspace_list_blurting(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.blurting_exercises.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(None)
    return docs


# List all diagram exercises for current user
@api_router.get("/workspace/diagrams")
async def workspace_list_diagrams(authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    docs = await db.diagram_exercises.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(None)
    return docs


# Delete a blurting exercise
@api_router.delete("/workspace/blurting/{exercise_id}")
async def workspace_delete_blurting(exercise_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    result = await db.blurting_exercises.delete_one({"id": exercise_id, "user_id": user["id"]})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Exercise not found")
    return {"ok": True}


# Delete a diagram exercise
@api_router.delete("/workspace/diagrams/{exercise_id}")
async def workspace_delete_diagram(exercise_id: str, authorization: Optional[str] = Header(None)):
    user = await auth_module.get_current_user(authorization)
    result = await db.diagram_exercises.delete_one({"id": exercise_id, "user_id": user["id"]})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Exercise not found")
    return {"ok": True}


# ---------- MOUNT ROUTER + MIDDLEWARE ----------
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------- STARTUP / SHUTDOWN ----------
@app.on_event("startup")
async def startup_db_client():
    existing = await auth_module.get_user_by_username("oscar")
    if not existing:
        await auth_module.create_user("oscar", "admin1234", is_admin=True)
        logger.info("Admin account created: oscar / admin1234 - CHANGE THIS PASSWORD!")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()

logger.info("Revisia API ready, model: %s, has_key: %s", DEFAULT_AI_MODEL, bool(OPENAI_API_KEY))
